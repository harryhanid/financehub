import io
import openpyxl
from modules.sponsorship.service import seed_perusahaan, create_transaksi, get_sponsorship_outstanding
from modules.sponsorship.exports import export_outstanding_excel


def test_sponsorship_outstanding_order(app):
    SMT = 1
    # Seed perusahaan with different cat2
    seed_perusahaan([
        ("CO_PROP", "Property", "Property"),
        ("CO_MINING", "Mining", "Mining"),
        ("CO_FIN", "Finance", "Finance"),
        ("CO_AGRI", "AGRI", "AGRI"),
        ("CO_APP", "APP", "APP"),
        ("CO_ENERGY", "ENERGY", "ENERGY"),
    ])

    categories = [
        ("CO_PROP", 1000000),
        ("CO_MINING", 2000000),
        ("CO_FIN", 3000000),
        ("CO_AGRI", 4000000),
        ("CO_APP", 5000000),
        ("CO_ENERGY", 6000000),
    ]
    for cat1, amt in categories:
        res = create_transaksi(SMT, {
            "cat": "Sponsorship",
            "description": "Sosial",
            "cat1": cat1,
            "tanggal_form_sponsor": "2026-02-01",
            "lembaga_perorangan": f"Lembaga {cat1}",
            "amount_request": amt,
        })
        assert res["ok"] is True

    sections, grand_total = get_sponsorship_outstanding(SMT)
    labels = [s["cat2_label"] for s in sections]
    
    # Standard order: APP, AGRI, Finance, Mining, Property, then others
    expected_order = [
        "Charging to APP",
        "Charging to AGRI",
        "Charging to Finance",
        "Charging to Mining",
        "Charging to Property",
        "Charging to ENERGY",
    ]
    assert labels == expected_order
    assert grand_total == sum(amt for _, amt in categories)


def test_export_outstanding_excel(app):
    SMT = 1
    seed_perusahaan([("CO_FIN2", "Finance", "Finance")])
    res = create_transaksi(SMT, {
        "cat": "Sponsorship",
        "description": "Sosial",
        "cat1": "CO_FIN2",
        "tanggal_form_sponsor": "2026-03-01",
        "lembaga_perorangan": "Lembaga Test Finance",
        "perihal": "Sponsorship Perihal",
        "amount_request": 7500000,
        "status": "Approved",
    })
    assert res["ok"] is True

    xls_bytes = export_outstanding_excel(SMT, {})
    assert len(xls_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    assert ws.title == "Outstanding Sponsorship"

    headers = [cell.value for cell in ws[1]]
    assert headers == ["NO", "Tanggal Approval", "Reff", "LEMBAGA/ PERORANGAN", "PERIHAL", "Amount in Rp", "Status Sponsor"]
    
    # Grand total row
    assert ws.cell(row=2, column=1).value == "Total Outstanding Sponsorship"
    assert ws.cell(row=2, column=6).value >= 7500000


def test_outstanding_export_route(client):
    from tests.test_sponsorship_routes import login, _select_smt
    login(client)
    _select_smt(client)

    resp = client.get("/sponsorship/outstanding/export")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert b"attachment; filename=Sponsorship_Outstanding.xlsx" in resp.headers["Content-Disposition"].encode("utf-8")
