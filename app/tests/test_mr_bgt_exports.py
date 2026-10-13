import io
from datetime import datetime

from openpyxl import load_workbook

from modules.management_report.bgt_report import build_bgt_report
from modules.management_report.exports import build_cc_workbook, build_zip
from tests.test_mr_bgt_report import ROWS

REPORT = {"cc": "AAA", "has_data": True, "groups": [], "grand": {}, "title": "unused"}


def _bgt(rows=ROWS, mm=8):
    bgt = build_bgt_report(rows, None, mm)
    bgt.update(title="President Office & ETF – Alpha Unit", cc_code="AAA",
               cc_desc="Alpha Unit", company_code="ETF")
    return bgt


def _ws(rows=ROWS, mm=8):
    wb = load_workbook(io.BytesIO(build_cc_workbook(REPORT, {}, None, None, _bgt(rows, mm))))
    assert wb.sheetnames == ["BGT Report per CC"]
    return wb["BGT Report per CC"]


def _fill(c):
    return c.fill.fgColor.rgb[-6:]


def test_header_block():
    ws = _ws()
    assert ws["C1"].value == "President Office & ETF – Alpha Unit" and ws["C2"].value == "In Rp Mio"
    assert ws["A2"].value == "AAA" and ws["A3"].value == "ETF - Alpha Unit"
    assert ws["C4"].value == "Period:" and ws["D4"].value == datetime(2026, 8, 31)
    assert ws["D4"].number_format == "mmm-yy"
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"C1:O1", "C2:O2", "C6:C8", "D6:G6", "I6:K6", "M6:N7", "O6:O8", "J7:K7"} <= merged
    assert ws["C6"].value == "Description" and ws["D6"].value == "Consolidation Budget"
    assert ws["I6"].value == "Consolidation Actual" and ws["O6"].value == "Remarks"
    assert ws["M6"].value == "Available Bgt YTD 2026 - Actual Aug 2026"
    assert ws["J7"].value == "Var A Aug 2026 to B Aug 2026"
    assert ws["G5"].value == "4\n( 2 + 3 )" and ws["K5"].value == "( 5 - 4 ) / 4"
    assert ws["E8"].value == datetime(2026, 8, 31) and ws["J8"].value == "Amount" and ws["K8"].value == "%"
    assert _fill(ws["C6"]) == "538DD5" and ws["C6"].font.color.rgb[-6:] == "FFFFFF"
    assert ws.sheet_view.zoomScale == 70
    assert ws.column_dimensions["C"].width == 59.3 and ws.column_dimensions["H"].width == 1.4


def test_body_rows_and_values():
    ws = _ws()
    assert ws["C10"].value == "OPEX" and _fill(ws["C10"]) == "95B3D7"
    assert [ws[f"{c}10"].value for c in "DEFGIJ"] == [19, 13, 1, 14, 4, -10]
    assert ws["K10"].value == -10 / 14 and ws["K10"].number_format == "0.0%"
    assert ws["M10"].value == 15 and ws["N10"].value == 15 / 19
    assert ws["C11"].value == "Rental, Tax" and _fill(ws["C11"]) == "CCECFF" and ws["C11"].font.bold
    assert ws["D11"].value == 16 and ws["I11"].value == 4
    assert ws["B12"].value == 70105120 and ws["C12"].value == "- Bgt 2026 : Rent Month : 1"
    assert (ws["D12"].value, ws["E12"].value, ws["F12"].value) == (10, 10, None)
    assert ws["C14"].value == "- Act 2026 : Rent Month : 1" and ws["I14"].value == 4
    assert ws["C18"].value == "CAPEX" and ws["C19"].value == "Office Furniture & Fixture"


def test_details_are_grouped_and_hidden_headers_are_not():
    ws = _ws()
    for r in (12, 13, 14, 16, 17, 20, 21):
        dim = ws.row_dimensions[r]
        assert dim.hidden and dim.outlineLevel == 2, r
    for r in (10, 11, 15, 18, 19):
        assert not ws.row_dimensions[r].hidden and not ws.row_dimensions[r].outlineLevel, r
    assert ws.sheet_properties.outlinePr.summaryBelow is False


def test_no_data_writes_only_the_header():
    wb = load_workbook(io.BytesIO(build_cc_workbook(
        REPORT, {}, None, None, {**_bgt([]), "has_data": False})))
    ws = wb["BGT Report per CC"]
    assert ws["C6"].value == "Description" and ws["C10"].value is None
    assert ws["D4"].value is None


def test_text_starting_with_equals_stays_text():
    rows = [dict(ROWS[0], category="01 - =SUM(1)", description="=cmd")]
    ws = _ws(rows)
    assert ws["C11"].value == "=SUM(1)" and ws["C11"].data_type == "s"
    assert ws["C12"].data_type == "s"


def test_legacy_call_keeps_the_old_report_sheet():
    wb = load_workbook(io.BytesIO(build_cc_workbook(REPORT)))
    assert wb.sheetnames == ["Report"]


def test_zip_uses_the_new_sheet_per_cc():
    import zipfile
    data = build_zip([REPORT], {}, None, None, {"AAA": _bgt()})
    zf = zipfile.ZipFile(io.BytesIO(data))
    assert load_workbook(io.BytesIO(zf.read("AAA.xlsx"))).sheetnames == ["BGT Report per CC"]
