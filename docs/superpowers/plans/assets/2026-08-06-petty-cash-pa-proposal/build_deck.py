"""Builds the Petty Cash & PA digital proposal deck (content-only, no design)."""
from pptx import Presentation
from pptx.util import Pt


def _add_slide(prs, title, bullets):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    for bullet in bullets:
        text, level = bullet if isinstance(bullet, tuple) else (bullet, 0)
        paragraph = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        paragraph.text = text
        paragraph.level = level
        paragraph.font.size = Pt(18 if level == 0 else 15)
    return slide


def build(output_path):
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = (
        "Digitalisasi Petty Cash & PA di Bawah Rp 10 Juta via PR Portal"
    )
    title_slide.placeholders[1].text = (
        "Proposal untuk [Nama Atasan/Manager] — [Nama Pengusul] — 6 Agustus 2026"
    )

    _add_slide(
        prs,
        "The Numbers: Volume vs Nilai",
        [
            "1.107 transaksi Petty Cash & PA (scope: SMT), total nilai Rp 74.262.296.882",
            ("741 transaksi (66,9%) bernilai < Rp 10 juta", 1),
            ("Nilai transaksi <10jt itu cuma Rp 1.750.525.855 — 2,36% dari total nilai", 1),
            "Insight: 67% volume transaksi hanya menyumbang 2,4% risiko finansial",
            "[Desain: big-number callout + chart donut/bar volume vs value]",
        ],
    )

    _add_slide(
        prs,
        "Masalah Proses Saat Ini (Manual Kertas/TTD Basah)",
        [
            "Lambat & bottleneck — approval berjenjang butuh approver fisik hadir, padahal nilainya kecil",
            "Tidak ada audit trail — dokumen fisik rawan hilang/rusak, sulit ditelusuri status & histori",
            "Rekonsiliasi manual ke Finance — data tidak terpusat, Finance input ulang dari kertas",
        ],
    )

    _add_slide(
        prs,
        "Breakdown Data",
        [
            "PA: 65,0% dari 881 transaksi berada di bawah Rp 10 juta",
            "Petty Cash: 74,3% dari 226 transaksi berada di bawah Rp 10 juta",
            "Top requester (<10jt): Chaterine (102), Grace (97), Caca (62), Rhiska (58), Tika (47), Suriyadi (44), Eli (40)",
            "Ini rutinitas harian banyak orang, bukan kasus minor/tepi",
        ],
    )

    _add_slide(
        prs,
        "Usulan: Digitalisasi via PR Portal",
        [
            "Petty Cash & PA < Rp 10 juta (scope SMT) dipindah jadi modul/alur baru di PR Portal",
            "Alur approval TETAP SAMA dengan PR Portal existing:",
            (
                "Requester -> Approver 1 -> Approver 2 (kondisional) -> Fincon (TF path) "
                "-> Verificator -> Releaser",
                1,
            ),
            "Tidak mengubah tata kelola/kontrol yang sudah berjalan — hanya mengubah media dari kertas ke digital",
        ],
    )

    _add_slide(
        prs,
        "Alur Baru: Before vs After",
        [
            "SEBELUM (manual): tanda tangan basah fisik, tracking manual, approver harus hadir",
            "SESUDAH (digital PR Portal): submit online, approval tercatat sistem, notifikasi otomatis",
            (
                "TAMBAHAN KONTROL: Dual-release wajib — 2 Releaser harus sama-sama approve "
                "sebelum dana cair (segregation of duties)",
                0,
            ),
            "[Desain: diagram 2 kolom Before | After dengan panah tahapan]",
        ],
    )

    _add_slide(
        prs,
        "Manfaat",
        [
            "Kecepatan proses — approval tidak tergantung kehadiran fisik approver",
            "Audit trail lengkap — histori tercatat, tidak ada dokumen hilang",
            "Rekonsiliasi otomatis — data langsung sync ke Supabase, tidak perlu input ulang Finance",
            "1 sistem terpadu dengan modul PR/PDO/PTA yang sudah berjalan",
        ],
    )

    _add_slide(
        prs,
        "Risiko & Mitigasi",
        [
            "Risiko: perubahan kebiasaan (change management) & reliance ke sistem",
            ("Mitigasi: masa transisi paralel (kertas + digital), training singkat", 1),
            "Dual-release justru MEMPERKUAT kontrol pencairan dibanding proses lama",
        ],
    )

    _add_slide(
        prs,
        "Next Steps & Ask",
        [
            "1. Approval proposal ini",
            "2. Technical spec & build modul di PR Portal",
            (
                "Catatan: perlu rekonsiliasi dengan spec Modul5 Payment Application yang "
                "sudah ada (item-based, beda scope)",
                1,
            ),
            "3. UAT dengan sample data",
            "4. Rollout SMT",
            "ASK: persetujuan untuk lanjut ke tahap spec teknis",
        ],
    )

    prs.save(output_path)


if __name__ == "__main__":
    build("petty-cash-pa-digital-proposal.pptx")
