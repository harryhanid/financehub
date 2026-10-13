"""Sheet 'BGT Report per CC' of the per-CC export, in the layout of the reference workbook."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.properties import Outline

_FONT = "Aptos Narrow"
_BLUE = PatternFill("solid", fgColor="538DD5")
_LEVEL = PatternFill("solid", fgColor="95B3D7")
_CATEGORY = PatternFill("solid", fgColor="CCECFF")
_NAME = PatternFill("solid", fgColor="F6C6AD")
_PERIOD = PatternFill("solid", fgColor="C1E4F5")
_NUM = r'_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'
_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_WIDTHS = {"A": 8.4, "B": 11.7, "C": 59.3, "D": 11.5, "E": 10.9, "F": 12.8, "G": 19.9, "H": 1.4,
           "I": 11.4, "J": 10.8, "K": 11.7, "L": 1.3, "M": 11.2, "N": 11.7, "O": 16}
_METRICS = (("d", "D"), ("e", "E"), ("f", "F"), ("g", "G"), ("i", "I"),
            ("j", "J"), ("k", "K"), ("m", "M"), ("n", "N"))
_FIRST_ROW = 10
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _font(size=11, bold=False, color=None):
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _text(cell, value):
    """Plain text: a value that starts with '=' must not become a formula."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _header(ws, bgt):
    period, year4 = bgt.get("period"), bgt.get("year4")
    mon = f"{_MONTHS[period - 1]} {year4}" if period and year4 else ""
    pdate = bgt.get("period_date")

    for r, h in ((1, 26), (2, 26), (3, 26), (4, 16), (5, 32), (6, 16), (7, 43.5), (8, 16)):
        ws.row_dimensions[r].height = h
    ws.merge_cells("C1:O1")
    ws.merge_cells("C2:O2")
    for ref, value in (("C1", bgt.get("title")), ("C2", "In Rp Mio")):
        _text(ws[ref], value)
        ws[ref].font = _font(20, True)
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    for ref, value in (("A2", bgt.get("cc_code")),
                       ("A3", f"{bgt.get('company_code')} - {bgt.get('cc_desc')}")):
        _text(ws[ref], value)
    for ref in ("A2", "B2", "A3", "B3"):
        ws[ref].font = _font(11, True)
        ws[ref].fill = _NAME

    ws["C4"], ws["D4"] = "Period:", pdate
    for ref in ("C4", "D4"):
        ws[ref].font = _font(12, True)
        ws[ref].fill = _PERIOD
    ws["C4"].alignment = Alignment(horizontal="left")
    ws["D4"].alignment = Alignment(horizontal="center")
    ws["D4"].number_format = "mmm-yy"

    for ref, value in (("D5", 1), ("E5", 2), ("F5", 3), ("G5", "4\n( 2 + 3 )"), ("I5", 5),
                       ("J5", "( 5 - 4 )"), ("K5", "( 5 - 4 ) / 4"), ("M5", "( 1 - 5 )"),
                       ("N5", "( 1 - 5 ) / 1")):
        ws[ref] = value
        ws[ref].font = _font(12, True)
        ws[ref].alignment = _CENTER

    for col in "CDEFGIJKMNO":
        for r in (6, 7, 8):
            c = ws[f"{col}{r}"]
            c.fill, c.font, c.alignment = _BLUE, _font(11, True, "FFFFFF"), _CENTER
    labels = {"C6": "Description", "D6": "Consolidation Budget", "I6": "Consolidation Actual",
              "M6": f"Available Bgt YTD {year4 or ''} - Actual {mon}", "O6": "Remarks",
              "D7": "Budget Full Year", "E7": "Budget YTD", "F7": "PTA YTD",
              "G7": "Budget Available YTD\n(Budget + PTA)", "I7": "Actual YTD",
              "J7": f"Var A {mon} to B {mon}", "J8": "Amount", "K8": "%", "M8": "Amount", "N8": "%"}
    for ref, value in labels.items():
        ws[ref] = value
    for col in "EFGI":
        ws[f"{col}8"] = pdate
        ws[f"{col}8"].number_format = "mmm-yy"
    for rng in ("C6:C8", "D6:G6", "I6:K6", "M6:N7", "O6:O8", "J7:K7"):
        ws.merge_cells(rng)


def _level_row(ws, r, label, m, fill, pct_format):
    for col in "CDEFGIJKMNO":
        ws[f"{col}{r}"].fill = fill
        ws[f"{col}{r}"].font = _font(11, True)
    _text(ws[f"C{r}"], label)
    for key, col in _METRICS:
        c = ws[f"{col}{r}"]
        c.value = m[key]
        c.number_format = pct_format if key in ("k", "n") else _NUM


def _detail_row(ws, r, d):
    coa = d["coa"]
    b = ws.cell(r, 2)
    b.value = int(coa) if str(coa).isdigit() else coa
    if isinstance(b.value, str):
        b.data_type = "s"
    b.font = _font()
    c = ws.cell(r, 3)
    _text(c, d["text"])
    c.font = _font()
    for key, col in (("d", "D"), ("e", "E"), ("f", "F"), ("i", "I")):
        if d[key] is not None:
            cell = ws[f"{col}{r}"]
            cell.value = d[key]
            cell.number_format = _NUM
            cell.font = _font()
    dim = ws.row_dimensions[r]
    dim.outlineLevel = 2
    dim.hidden = True


def write_bgt_sheet(ws, bgt: dict) -> None:
    ws.sheet_view.zoomScale = 70
    ws.sheet_format.outlineLevelRow = 2
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True)
    for col, width in _WIDTHS.items():
        ws.column_dimensions[col].width = width
    _header(ws, bgt)
    r = _FIRST_ROW
    for level in bgt.get("levels", []):
        _level_row(ws, r, level["label"], level["m"], _LEVEL, "0.0%")
        r += 1
        for cat in level["categories"]:
            _level_row(ws, r, cat["label"], cat["m"], _CATEGORY, "0%")
            r += 1
            for d in cat["details"]:
                _detail_row(ws, r, d)
                r += 1
