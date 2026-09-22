import sys
import io

def patch():
    with open(r'C:\Financehub\app\modules\sahabat_etf\report_xlsx.py', 'r', encoding='utf-8') as f:
        content = f.read()

    old_func = '''def _write_title(ws, row):
    ws.merge_cells(f"B{row}:N{row}")
    _set(ws, f"B{row}", f"REPORT PER {_today_title()}", font=_FONT_TITLE, align=_ALIGN_LEFT)
    _set(ws, f"O{row}", "(Rp, Dalam Jutaan)", font=_FONT_UNIT, align=_ALIGN_RIGHT)
    return row + 2'''

    new_func = '''def _write_title(ws, row, report_year, report_month):
    now = datetime.now()
    month_name = _INDO_MONTHS[report_month].upper() if report_month in _INDO_MONTHS else ""
    title_text = f"REPORT PER {month_name} {report_year}, dicetak {now.day}/{now.month}/{now.year}"
    
    ws.merge_cells(f"B{row}:N{row}")
    _set(ws, f"B{row}", title_text, font=_FONT_TITLE, align=_ALIGN_LEFT)
    _set(ws, f"O{row}", "(Rp, Dalam Jutaan)", font=_FONT_UNIT, align=_ALIGN_RIGHT)
    return row + 2'''

    content = content.replace(old_func, new_func)

    old_call = '''row = _write_title(ws, 1)'''
    new_call = '''row = _write_title(ws, 1, data["report_year"], data.get("report_month", 12))'''
    content = content.replace(old_call, new_call)

    with open(r'C:\Financehub\app\modules\sahabat_etf\report_xlsx.py', 'w', encoding='utf-8') as f:
        f.write(content)

    code = '''

def build_detail_workbook(rows: list) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    import io

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Detail Transaksi"
    
    headers = ["Sumber", "Kode Siswa", "Nama", "Tanggal", "Tgl Final", "Kategori 1", "Kategori 2", "Pillar", "Amount", "Status"]
    sheet.append(headers)
    
    for i in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=i)
        cell.font = Font(bold=True)
        try:
            cell.fill = _FILL_NAVY
            cell.font = Font(color=_WHITE, bold=True)
        except NameError:
            pass
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for r in rows:
        sheet.append([
            r.get("sumber", ""),
            r.get("siswa_code", ""),
            r.get("nama", ""),
            r.get("tanggal", ""),
            r.get("tgl_final", ""),
            r.get("cat1", ""),
            r.get("cat2", ""),
            r.get("pillar", ""),
            r.get("amount", 0),
            r.get("status", "")
        ])
        
    import openpyxl.utils
    for i, _ in enumerate(headers, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 15

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
'''

    with open(r'C:\Financehub\app\modules\sahabat_etf\report_xlsx.py', 'a', encoding='utf-8') as f:
        f.write(code)

if __name__ == '__main__':
    patch()
