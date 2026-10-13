import io
import zipfile

import pytest
from openpyxl import load_workbook

from modules.management_report import service as svc
from modules.management_report import ho_report as hr
from modules.management_report.exports import (APPEND_HEADERS, build_append_workbook,
                                                build_cc_workbook, build_zip, safe_filename)
from tests.test_mr_routes import login
from tests.test_mr_service import mkrow


@pytest.fixture
def seeded():
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="TFKHAR", category="01 - Rental", mm=3, mm1=3, description="Rent", bgt=3.2, act=0),
        mkrow(cc="TFKHAR", category="06 - Gift", mm=3, mm1=3, description="Gift", bgt=33.5, act=7.0),
        mkrow(cc="POSKHU - AR", category="01 - Rental", description="x", bgt=2.0),
        mkrow(cc="EMPTYCC", category="01 - Rental", mm=12, description="y", bgt=1.0),
    ])


def test_safe_filename():
    assert safe_filename('A/B:C*D?"E<F>G|H\\I') == "A_B_C_D__E_F_G_H_I"
    assert safe_filename("POSKHU - AR") == "POSKHU - AR"
    assert safe_filename("  ") == "CC"


def test_workbook_layout_and_values(seeded):
    rep = svc.build_report(2, "TFKHAR")
    ws = load_workbook(io.BytesIO(build_cc_workbook(rep, {"yy": None, "mm": None, "coas": []}))).active
    assert ws["A1"].value == "TFKHAR" and ws["A2"].value == "In Rp Mio"
    assert [ws.cell(4 + i, 1).value for i in range(4)] == ["CC", "YY", "MM", "COA"]
    assert [ws.cell(4 + i, 2).value for i in range(4)] == ["TFKHAR", "All", "All", "All"]
    assert [c.value for c in ws[10][:11]] == [
        "Category", "Custom", "MM.1", "Description", "Activity", "BGT", "PTA",
        "BGT_Available", "ACT", "Diff_ACT-BGT", "Diff Percentage"]
    vals = [[c.value for c in row[:11]] for row in ws.iter_rows(min_row=11, max_row=15)]
    assert vals[0][0] == "01 - Rental" and vals[0][3] == "Rent" and vals[0][5] == 3.2
    assert vals[1][0] == "01 - Rental Total" and vals[1][7] == 3.2 and vals[1][10] == -1.0
    assert vals[2][0] == "06 - Gift" and vals[3][0] == "06 - Gift Total"
    assert vals[3][8] == 7.0 and vals[3][10] == pytest.approx(-26.5 / 33.5, abs=1e-6)
    g = vals[4]
    assert g[0] == "Grand Total" and g[5] == pytest.approx(36.7) and g[8] == 7.0
    assert ws["F11"].number_format == "#,##0.0" and ws["K12"].number_format == "0.0%"
    assert ws.freeze_panes == "A11"


def test_blank_percentage_is_empty_cell(seeded):
    svc.save_upload(2, "b.xlsb", "u", [mkrow(cc="Z", bgt=0, pta=0, act=3)])
    ws = load_workbook(io.BytesIO(build_cc_workbook(svc.build_report(2, "Z")))).active
    assert ws["K11"].value is None


def test_zip_contains_one_file_per_cc(seeded):
    reps = [svc.build_report(2, cc) for cc in ("TFKHAR", "POSKHU - AR")]
    zf = zipfile.ZipFile(io.BytesIO(build_zip(reps)))
    assert sorted(zf.namelist()) == ["POSKHU - AR.xlsx", "TFKHAR.xlsx"]
    ws = load_workbook(io.BytesIO(zf.read("TFKHAR.xlsx"))).active
    assert ws["A1"].value == "TFKHAR"


def test_export_route_downloads_xlsx(client, seeded):
    login(client)
    resp = client.get("/management-report/export?cc=TFKHAR")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "TFKHAR.xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["BGT Report per CC", "BGT HO Report per CC", "Append"]
    assert wb["BGT Report per CC"]["C1"].value.endswith("TFKHAR")


def test_export_route_honours_filters_and_404(client, seeded):
    login(client)
    assert client.get("/management-report/export?cc=NOPE").status_code == 404
    assert client.get("/management-report/export").status_code == 400
    assert client.get("/management-report/export?cc=EMPTYCC&mm=1").status_code == 404


def test_export_all_skips_empty_ccs(client, seeded):
    login(client)
    resp = client.get("/management-report/export-all?mm=3")
    assert resp.status_code == 200 and resp.mimetype == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(resp.data)).namelist()
    assert "EMPTYCC.xlsx" not in names and "TFKHAR.xlsx" in names
    # mm=11 is year to date: TFKHAR/POSKHU have lines up to then, so they export now
    assert client.get("/management-report/export-all?mm=11").status_code == 200
    assert client.get("/management-report/export-all?yy=1999").status_code == 404


def test_export_requires_login(client):
    assert client.get("/management-report/export?cc=X").status_code == 302
    assert client.get("/management-report/export-all").status_code == 302


@pytest.fixture
def ho_seeded():
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", category="01 - Rental", coa="70105120", bgt=3.2, act=0.5, description="Rent"),
        mkrow(cc="AAA", category="01 - Rental", coa="71000110", act=-97.5, description="Interest"),
        mkrow(cc="AAA", category="01 - Rental", coa="70102114", bgt=12.0, act=5.6, description="BPJS"),
    ], meta={"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"})


def _row_with(ws, value, col):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == value)


def _ho_workbook():
    rep = svc.build_report(2, "AAA")
    rep["title"] = "President Office & ETF – Alpha Unit"
    return load_workbook(io.BytesIO(build_cc_workbook(rep, {}, hr.build_ho_report(2, "AAA"))))


def test_workbook_gets_a_second_sheet_with_the_ho_report(ho_seeded):
    wb = _ho_workbook()
    assert wb.sheetnames == ["Report", "BGT HO Report per CC"]
    assert wb["Report"]["A1"].value == "President Office & ETF – Alpha Unit"
    ws = wb["BGT HO Report per CC"]
    assert ws["A1"].value == "President Office & ETF" and ws["A2"].value.startswith("Head Office")
    assert ws["A4"].value == "Division/Dept.: Alpha Unit"
    r = _row_with(ws, 70105120, 1)
    assert ws.cell(r, 1).number_format == "000-00-000"
    assert ws.cell(r, 2).value == "Equipment Rent Expense - 3rd Party"
    assert (ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value) == (3200.0, 0.0, 500.0)
    assert ws.cell(r, 4).number_format.startswith("_(* #,##0_)")
    t = _row_with(ws, "Total - 701-05-xxx", 1)
    assert ws.cell(t, 4).value == 3200.0 and ws.cell(t, 1).font.bold is True
    g = _row_with(ws, "Grand Total included Depreciation", 1)
    assert ws.cell(g, 4).value == 15200.0                    # 3,200 + 12,000 (70102114 counts in Opex)
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert texts.count("s/d Aug") >= 9                        # 3 sections x 3 columns
    summary_total = _row_with(ws, "Total", 2)
    assert ws.cell(summary_total, 4).value == 3200.0         # 70102114 is in no summary line


def test_ho_sheet_footnotes_for_unmapped_coa_and_summary_gap(ho_seeded):
    ws = _ho_workbook()["BGT HO Report per CC"]
    g = _row_with(ws, "Grand Total included Depreciation", 1)
    assert str(ws.cell(g + 1, 1).value).startswith("Tidak termasuk COA 71000110")
    assert "ACT -97.500" in ws.cell(g + 1, 1).value
    s = _row_with(ws, "Total", 2)
    note = str(ws.cell(s + 1, 1).value)
    assert note.startswith("Tidak ada di ringkasan") and "70102114" in note


def test_zip_has_both_sheets_per_cc(ho_seeded):
    reps = [svc.build_report(2, "AAA")]
    zf = zipfile.ZipFile(io.BytesIO(build_zip(reps, {}, {"AAA": hr.build_ho_report(2, "AAA")})))
    assert zf.namelist() == ["AAA.xlsx"]
    assert load_workbook(io.BytesIO(zf.read("AAA.xlsx"))).sheetnames == ["Report", "BGT HO Report per CC"]


def test_workbook_without_ho_data_keeps_a_single_sheet(ho_seeded):
    wb = load_workbook(io.BytesIO(build_cc_workbook(svc.build_report(2, "AAA"))))
    assert wb.sheetnames == ["Report"]


def test_export_all_zip_has_both_sheets_in_every_workbook(client, seeded):
    login(client)
    resp = client.get("/management-report/export-all?mm=3")
    zf = zipfile.ZipFile(io.BytesIO(resp.data))
    assert zf.namelist()
    for name in zf.namelist():
        assert load_workbook(io.BytesIO(zf.read(name))).sheetnames == ["BGT Report per CC", "BGT HO Report per CC", "Append"]


def _append_rows(cc, **kw):
    return svc.query_append(2, cc, **kw)["rows"]


def test_workbook_gets_an_append_sheet_with_the_raw_rows(seeded):
    rep = svc.build_report(2, "TFKHAR")
    rows = _append_rows("TFKHAR")
    wb = load_workbook(io.BytesIO(build_cc_workbook(rep, {}, None, rows)))
    assert wb.sheetnames == ["Report", "Append"]
    ws = wb["Append"]
    assert [c.value for c in ws[1]] == APPEND_HEADERS
    assert ws.max_row == 1 + len(rows) == 3
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:L3"
    bgt = APPEND_HEADERS.index("BGT") + 1
    assert sum(ws.cell(r, bgt).value for r in range(2, ws.max_row + 1)) == \
        pytest.approx(rep["grand"]["bgt"])                      # reconciles with sheet Report
    assert ws.cell(2, 1).value == 70105120                      # COA as a number
    assert ws.cell(2, bgt).number_format == "#,##0.0##"
    assert ws.cell(2, APPEND_HEADERS.index("Description") + 1).value == "Rent"


def test_append_sheet_comes_after_the_ho_sheet(ho_seeded):
    rep = svc.build_report(2, "AAA")
    wb = load_workbook(io.BytesIO(
        build_cc_workbook(rep, {}, hr.build_ho_report(2, "AAA"), _append_rows("AAA"))))
    assert wb.sheetnames == ["Report", "BGT HO Report per CC", "Append"]
    assert wb["Append"].max_row == 1 + 3


def test_append_sheet_values_and_empty_cells(seeded):
    rows = _append_rows("TFKHAR")[:1]
    rows[0].update(coa="70105120A", mm=None, description=None)
    ws = load_workbook(io.BytesIO(build_cc_workbook(svc.build_report(2, "TFKHAR"), {}, None, rows)))["Append"]
    assert ws.cell(2, 1).value == "70105120A"                   # not all digits: kept as text
    assert ws.cell(2, 2).value is None                          # MM empty
    assert ws.cell(2, APPEND_HEADERS.index("Description") + 1).value is None


def test_workbook_without_append_rows_has_no_append_sheet(seeded):
    wb = load_workbook(io.BytesIO(build_cc_workbook(svc.build_report(2, "TFKHAR"))))
    assert "Append" not in wb.sheetnames


def test_zip_gives_every_cc_its_own_append_rows(seeded):
    reps = [svc.build_report(2, cc) for cc in ("TFKHAR", "POSKHU - AR")]
    by_cc = {cc: _append_rows(cc) for cc in ("TFKHAR", "POSKHU - AR")}
    zf = zipfile.ZipFile(io.BytesIO(build_zip(reps, {}, None, by_cc)))
    counts = {}
    for name in zf.namelist():
        wb = load_workbook(io.BytesIO(zf.read(name)))
        assert wb.sheetnames == ["Report", "Append"]
        counts[name] = wb["Append"].max_row - 1
    assert counts == {"TFKHAR.xlsx": 2, "POSKHU - AR.xlsx": 1}


def test_standalone_append_workbook_has_only_the_append_sheet(seeded):
    wb = load_workbook(io.BytesIO(build_append_workbook(_append_rows(None))))
    assert wb.sheetnames == ["Append"] and wb["Append"].max_row == 1 + 4


def test_ho_sheet_is_named_like_the_master_and_styled_like_the_reference(ho_seeded):
    wb = _ho_workbook()
    assert wb.sheetnames == ["Report", "BGT HO Report per CC"]
    ws = wb["BGT HO Report per CC"]
    assert ws["G2"].value.year == 2026 and ws["G2"].value.month == 8
    assert ws["G2"].fill.fgColor.rgb[-6:] == "7CEB99" and ws["G3"].value == "AAA"
    assert ws["G3"].fill.fgColor.rgb[-6:] == "7CEB99"
    head = _row_with(ws, "GL Account", 1)
    assert ws.cell(head, 1).fill.fgColor.rgb[-6:] == "FFCC99"
    assert ws.cell(head, 1).border.left.style == "medium" and ws.cell(head, 1).border.top.style == "medium"
    assert ws.cell(head + 1, 4).border.bottom.style == "medium"
    r = _row_with(ws, 70105120, 1)
    assert ws.cell(r, 4).fill.fgColor.rgb[-6:] == "CCFFFF"
    assert ws.cell(r, 4).border.left.style == "medium" and ws.cell(r, 4).border.right.style == "medium"
    assert ws.cell(r, 1).border.left.style == "medium" and ws.cell(r, 1).border.right.style == "thin"
    t = _row_with(ws, "Total - 701-05-xxx", 1)
    assert ws.cell(t, 1).border.top.style == "thin" and ws.cell(t, 4).border.bottom.style == "thin"
    assert ws.cell(t, 1).alignment.horizontal == "centerContinuous"
    g = _row_with(ws, "Grand Total included Depreciation", 1)
    assert ws.cell(g, 4).border.top.style == "medium"
    assert ws.column_dimensions["A"].width == 17.2 and ws.column_dimensions["B"].width == 52


S1, S2, S3 = "BGT Report per CC", "BGT HO Report per CC", "Append"


def _single_names(client, qs):
    r = client.get("/management-report/export?cc=TFKHAR" + qs)
    assert r.status_code == 200 and 'filename="TFKHAR.xlsx"' in r.headers["Content-Disposition"]
    return load_workbook(io.BytesIO(r.data)).sheetnames


def test_export_one_sheet_choice(client, seeded):
    login(client)
    assert _single_names(client, "&sheets=1") == [S1]
    assert _single_names(client, "&sheets=12") == [S1, S2]
    assert _single_names(client, "&sheets=123") == [S1, S2, S3]
    assert _single_names(client, "") == [S1, S2, S3]
    assert _single_names(client, "&sheets=bogus") == [S1, S2, S3]


def _zip(client, qs):
    r = client.get("/management-report/export-all?mm=3" + qs)
    assert r.status_code == 200 and r.mimetype == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(r.data))
    assert "TFKHAR.xlsx" in zf.namelist()
    names = {tuple(load_workbook(io.BytesIO(zf.read(n))).sheetnames) for n in zf.namelist()}
    return r.headers["Content-Disposition"], [list(n) for n in names]


def test_export_all_sheet_choice(client, seeded):
    login(client)
    cd, sheets = _zip(client, "&sheets=1")
    assert "sheet 1.zip" in cd and "sheet 1-2" not in cd and sheets == [[S1]]
    cd, sheets = _zip(client, "&sheets=12")
    assert "sheet 1-2.zip" in cd and sheets == [[S1, S2]]
    cd, sheets = _zip(client, "")
    assert "sheet " not in cd and sheets == [[S1, S2, S3]]
    cd, sheets = _zip(client, "&append=0")                   # legacy: no Append sheet
    assert "sheet 1-2.zip" in cd and sheets == [[S1, S2]]
    cd, sheets = _zip(client, "&sheets=1&append=0")          # sheets wins
    assert "sheet 1.zip" in cd and sheets == [[S1]]


def test_index_has_one_export_all_button_and_the_sheet_popover(client, seeded):
    login(client)
    html = client.get("/management-report/").data.decode()
    assert html.count('id="mr-export-all"') == 1
    assert "mr-export-all-lite" not in html and "tanpa Append" not in html
    assert 'id="mr-export-pop"' in html
    assert 'name="mr-export-sheets" value="1"' in html and 'name="mr-export-sheets" value="12"' in html
    assert 'name="mr-export-sheets" value="123" checked' in html
    assert 'id="mr-export-go"' in html and 'id="mr-export-cancel"' in html


@pytest.fixture
def cat_seeded():
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", category="01 - Rental", description="Rent", bgt=3.0),
        mkrow(cc="AAA", coa="70106100", category="02 - Transport", description="Taxi", bgt=5.0),
        mkrow(cc="BBB", coa="70106100", category="02 - Transport", description="Bus", bgt=7.0),
    ], meta={"cc_names": {"AAA": "Alpha", "BBB": "Beta"}, "period": "2026-08-31"})


def test_filter_options_list_the_categories(cat_seeded):
    assert svc.get_filter_options(2)["cats"] == ["01 - Rental", "02 - Transport"]
    assert svc.get_filter_options(99)["cats"] == []


def test_build_report_and_append_follow_the_category_filter(cat_seeded):
    rep = svc.build_report(2, "AAA", cats=["02 - Transport"])
    assert [g["category"] for g in rep["groups"]] == ["02 - Transport"]
    assert svc.build_report(2, "AAA", cats=["01 - Rental", "02 - Transport"])["grand"]["bgt"] == 8.0
    assert svc.build_report(2, "AAA")["grand"]["bgt"] == 8.0                # no filter = all
    assert svc.query_append(2, "AAA", cats=["01 - Rental"])["total"] == 1
    assert svc.list_ccs_with_data(2, cats=["01 - Rental"]) == ["AAA"]
    assert svc.list_ccs_with_data(2, cats=["02 - Transport"]) == ["AAA", "BBB"]


def test_export_follows_the_category_filter(client, cat_seeded):
    login(client)
    wb = load_workbook(io.BytesIO(client.get(
        "/management-report/export?cc=AAA&cat=02 - Transport").data))
    ws = wb["BGT Report per CC"]
    assert ws["C10"].value == "OPEX" and ws["D10"].value == 5
    assert wb["Append"].max_row == 1 + 1
    zf = zipfile.ZipFile(io.BytesIO(client.get(
        "/management-report/export-all?cat=01 - Rental&append=0").data))
    assert zf.namelist() == ["AAA.xlsx"]


def test_report_table_and_page_follow_the_category_filter(client, cat_seeded):
    login(client)
    html = client.get("/management-report/table?cc=AAA&cat=01 - Rental").data.decode()
    assert "Rent" in html and "Taxi" not in html
    page = client.get("/management-report/?cc=AAA&cat=01 - Rental").data.decode()
    assert 'id="mr-cat-panel"' in page and 'value="02 - Transport"' in page
    assert 'value="01 - Rental" checked' in page and 'value="02 - Transport" checked' not in page
    assert "cat=01" in page                                   # export links keep the filter
