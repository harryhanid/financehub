import os
import docx
from modules.sop.migrate import migrate_sop_documents
from modules.sop.service import list_documents


def _make_source_tree(root):
    modul_dir = os.path.join(root, "procurement")
    os.makedirs(modul_dir, exist_ok=True)
    document = docx.Document()
    document.add_heading("SOP Master Vendor", level=0)
    document.add_heading("1. Tujuan", level=1)
    document.add_paragraph("⚠ [PERLU VERIFIKASI: contoh]")
    document.save(os.path.join(modul_dir, "1-SOP_Master_Vendor.docx"))
    return root


def test_migrate_imports_docx_files_as_aktif(tmp_path):
    source_root = _make_source_tree(str(tmp_path))
    result = migrate_sop_documents(source_root)
    assert result["imported"] == 1

    docs = list_documents(modul="procurement")
    assert len(docs) == 1
    assert docs[0]["nama"] == "SOP Master Vendor"
    assert docs[0]["status"] == "aktif"
    assert docs[0]["perlu_verifikasi_count"] == 1
    assert docs[0]["kode"] == "1"


def test_migrate_seeds_out_of_scope_rows(tmp_path):
    source_root = _make_source_tree(str(tmp_path))
    result = migrate_sop_documents(source_root)
    assert result["out_of_scope"] == 8

    belum = list_documents(status="belum_dibuat")
    assert len(belum) == 8
    assert any(d["nama"] == "Penagihan Diplomacy & Advisory Fee" for d in belum)
    assert any(d["nama"] == "Payroll" for d in belum)
