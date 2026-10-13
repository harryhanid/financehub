"""Management Report exports: one xlsx per cost center, or a zip of all."""
import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from modules.management_report import bgt_exports, ho_exports

HEADERS = ["Category", "Custom", "MM.1", "Description", "Activity", "BGT", "PTA",
           "BGT_Available", "ACT", "Diff_ACT-BGT", "Diff Percentage"]
HEADER_ROW = 10
NUM_FMT = "#,##0.0"
PCT_FMT = "0.0%"
_FILL_HEAD = PatternFill("solid", fgColor="E2E8F0")
_FILL_SUB = PatternFill("solid", fgColor="F1F5F9")
_BOLD = Font(bold=True)
_WIDTHS = {"A": 38, "B": 8, "C": 6, "D": 70, "E": 24,
           "F": 14, "G": 14, "H": 14, "I": 14, "J": 14, "K": 14}

APPEND_HEADERS = ["COA", "MM", "YY", "CC", "Category", "Activity", "Description",
                  "PTA", "Custom", "MM.1", "ACT", "BGT"]
_APPEND_KEYS = ["coa", "mm", "yy", "cc", "category", "activity", "description",
                "pta", "custom", "mm1", "act", "bgt"]
_APPEND_WIDTHS = {"A": 11, "B": 5, "C": 5, "D": 14, "E": 36, "F": 24, "G": 70,
                  "H": 12, "I": 8, "J": 6, "K": 12, "L": 12}
_APPEND_NUM_FMT = "#,##0.0##"


def safe_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]', "_", name or "").strip()
    return cleaned or "CC"


def _put_metrics(ws, row: int, m: dict, bold: bool = False, fill=None) -> None:
    for col, key in zip("FGHIJ", ("bgt", "pta", "available", "act", "diff")):
        c = ws[f"{col}{row}"]
        c.value = m[key]
        c.number_format = NUM_FMT
    k = ws[f"K{row}"]
    k.value = m["pct"]
    k.number_format = PCT_FMT
    if bold or fill:
        for col in "ABCDEFGHIJK":
            c = ws[f"{col}{row}"]
            if bold:
                c.font = _BOLD
            if fill:
                c.fill = fill


def _append_value(key: str, value):
    if value is None:
        return None
    if key == "coa":
        return int(value) if str(value).isdigit() else value
    return value


def add_append_sheet(wb, rows: list, name: str = "Append"):
    """Raw Append rows as a plain table: header row, one row per Append line, values only."""
    ws = wb.create_sheet(name)
    for col, header in enumerate(APPEND_HEADERS, start=1):
        c = ws.cell(1, col, header)
        c.font = _BOLD
        c.fill = _FILL_HEAD
        c.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=2):
        for col, key in enumerate(_APPEND_KEYS, start=1):
            c = ws.cell(r, col, _append_value(key, row.get(key)))
            if key in ("pta", "act", "bgt"):
                c.number_format = _APPEND_NUM_FMT
    for col, width in _APPEND_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(len(rows) + 1, 2)}"
    return ws


def build_append_workbook(rows: list) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    add_append_sheet(wb, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_legacy_report(ws, report: dict, filters: dict) -> None:
    ws["A1"] = report.get("title") or report["cc"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "In Rp Mio"
    coas = filters.get("coas") or []
    for i, (label, value) in enumerate([
        ("CC", report["cc"]),
        ("YY", filters.get("yy") if filters.get("yy") is not None else "All"),
        ("MM", filters.get("mm") if filters.get("mm") is not None else "All"),
        ("COA", ", ".join(coas) if coas else "All"),
    ]):
        ws.cell(4 + i, 1, label).font = _BOLD
        ws.cell(4 + i, 2, value)

    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = _BOLD
        c.fill = _FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i > 5 else "left")

    r = HEADER_ROW + 1
    for g in report["groups"]:
        for idx, row in enumerate(g["rows"]):
            ws.cell(r, 1, g["category"] if idx == 0 else None)
            ws.cell(r, 2, row["custom"])
            ws.cell(r, 3, row["mm1"])
            ws.cell(r, 4, row["description"])
            ws.cell(r, 5, row["activity"])
            _put_metrics(ws, r, row)
            r += 1
        ws.cell(r, 1, f"{g['category']} Total")
        _put_metrics(ws, r, g["total"], bold=True, fill=_FILL_SUB)
        r += 1
    if report["groups"]:
        ws.cell(r, 1, "Grand Total")
        _put_metrics(ws, r, report["grand"], bold=True, fill=_FILL_HEAD)

    for col, w in _WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{HEADER_ROW + 1}"


def build_cc_workbook(report: dict, filters: dict = None, ho: dict = None, append_rows: list = None,
                      bgt: dict = None) -> bytes:
    filters = filters or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    if bgt is not None:
        ws.title = "BGT Report per CC"
        bgt_exports.write_bgt_sheet(ws, bgt)
    else:
        _write_legacy_report(ws, report, filters)
    if ho is not None:
        ho_exports.add_ho_sheet(wb, ho)
    if append_rows is not None:
        add_append_sheet(wb, append_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_zip(reports: list, filters: dict = None, ho_by_cc: dict = None, append_by_cc: dict = None,
              bgt_by_cc: dict = None) -> bytes:
    buf = io.BytesIO()
    used = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for rep in reports:
            name = safe_filename(rep["cc"])
            base, n = name, 2
            while name in used:          # two CCs can sanitise to the same name
                name, n = f"{base}_{n}", n + 1
            used.add(name)
            zf.writestr(f"{name}.xlsx",
                        build_cc_workbook(rep, filters, (ho_by_cc or {}).get(rep["cc"]),
                                          (append_by_cc or {}).get(rep["cc"]),
                                          (bgt_by_cc or {}).get(rep["cc"])))
    return buf.getvalue()
