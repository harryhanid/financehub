import os
import sys

sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
import build_deck

OUT = os.path.join(os.path.dirname(__file__), "_test_deck.pptx")

EXPECTED_TITLES = [
    "Digitalisasi Petty Cash & PA di Bawah Rp 10 Juta via PR Portal",
    "The Numbers: Volume vs Nilai",
    "Masalah Proses Saat Ini (Manual Kertas/TTD Basah)",
    "Breakdown Data",
    "Usulan: Digitalisasi via PR Portal",
    "Alur Baru: Before vs After",
    "Manfaat",
    "Risiko & Mitigasi",
    "Next Steps & Ask",
]


def _all_text(slide):
    chunks = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                chunks.append("".join(run.text for run in para.runs))
    return "\n".join(chunks)


def test_build_deck_has_nine_slides_with_expected_titles():
    build_deck.build(OUT)
    prs = Presentation(OUT)
    assert len(prs.slides) == 9
    for slide, expected_title in zip(prs.slides, EXPECTED_TITLES):
        assert slide.shapes.title.text == expected_title


def test_slide2_contains_key_figures():
    build_deck.build(OUT)
    prs = Presentation(OUT)
    text = _all_text(prs.slides[1])
    assert "741 transaksi (66,9%)" in text
    assert "Rp 1.750.525.855" in text
    assert "2,36%" in text


def test_slide9_contains_modul5_reconciliation_note():
    build_deck.build(OUT)
    prs = Presentation(OUT)
    text = _all_text(prs.slides[8])
    assert "Modul5 Payment Application" in text
