# app/tests/test_sop_routes.py
from modules.sop.service import create_document


def login(client, username="admin", password="Admin@123"):
    return client.post("/auth/login", json={"username": username, "password": password})


def test_index_requires_login(client):
    resp = client.get("/sop/")
    assert resp.status_code == 302


def test_index_renders_after_login(client):
    login(client)
    resp = client.get("/sop/")
    assert resp.status_code == 200


def test_index_lists_created_documents(client):
    login(client)
    create_document(modul="procurement", nama="SOP Master Vendor", kode="1")
    resp = client.get("/sop/")
    assert b"SOP Master Vendor" in resp.data


def test_index_filters_by_modul(client):
    login(client)
    create_document(modul="procurement", nama="SOP Hanya Procurement", kode="1")
    create_document(modul="finance-payment", nama="SOP Hanya Finance", kode="1")
    resp = client.get("/sop/?modul=procurement")
    assert b"SOP Hanya Procurement" in resp.data
    assert b"SOP Hanya Finance" not in resp.data


def test_index_search_by_nama(client):
    login(client)
    create_document(modul="it-operations", nama="SOP Zoom Meeting", kode="3")
    resp = client.get("/sop/?q=zoom")
    assert b"SOP Zoom Meeting" in resp.data


import os
import docx
import config


def _write_sample_docx(rel_path):
    full_path = os.path.join(config.SOP_STORAGE_DIR, rel_path)
    os.makedirs(os.path.dirname(full_path), exist_ok=True)
    document = docx.Document()
    document.add_heading("SOP Baca Uji", level=0)
    document.add_heading("1. Tujuan", level=1)
    document.add_paragraph("Tujuan uji baca.", style="List Bullet")
    document.save(full_path)
    return full_path


def test_read_page_renders_parsed_content(client):
    login(client)
    _write_sample_docx("procurement/1-uji.docx")
    doc_id = create_document(modul="procurement", nama="SOP Baca Uji", kode="1",
                              file_path="procurement/1-uji.docx")
    resp = client.get(f"/sop/{doc_id}")
    assert resp.status_code == 200
    assert b"Tujuan uji baca." in resp.data


def test_read_page_shows_placeholder_for_belum_dibuat(client):
    login(client)
    doc_id = create_document(modul="hr-personnel", nama="SOP Belum Ada",
                              status="belum_dibuat", catatan="Tidak ada sumber")
    resp = client.get(f"/sop/{doc_id}")
    assert resp.status_code == 200
    assert b"Tidak ada sumber" in resp.data


def test_read_page_redirects_for_unknown_id(client):
    login(client)
    resp = client.get("/sop/999999")
    assert resp.status_code == 302


def test_read_page_shows_placeholder_when_file_missing_on_disk(client):
    login(client)
    doc_id = create_document(modul="procurement", nama="SOP File Hilang", kode="2",
                              file_path="procurement/2-tidak-ada.docx")
    resp = client.get(f"/sop/{doc_id}")
    assert resp.status_code == 200
    assert b"File tidak ditemukan di storage" in resp.data


def test_download_returns_the_uploaded_file(client):
    login(client)
    _write_sample_docx("procurement/1-download-uji.docx")
    doc_id = create_document(modul="procurement", nama="SOP Download Uji", kode="1",
                              file_path="procurement/1-download-uji.docx")
    resp = client.get(f"/sop/{doc_id}/download")
    assert resp.status_code == 200
    assert resp.headers["Content-Disposition"].startswith("attachment")


def test_download_redirects_for_belum_dibuat(client):
    login(client)
    doc_id = create_document(modul="hr-personnel", nama="SOP Tanpa File", status="belum_dibuat")
    resp = client.get(f"/sop/{doc_id}/download")
    assert resp.status_code == 302


def test_download_redirects_when_file_missing_on_disk(client):
    login(client)
    doc_id = create_document(modul="procurement", nama="SOP Download Hilang", kode="3",
                              status="aktif", file_path="procurement/3-tidak-ada.docx")
    resp = client.get(f"/sop/{doc_id}/download")
    assert resp.status_code == 302


from io import BytesIO


def test_upload_replaces_file_and_records_history(client):
    login(client)
    _write_sample_docx("procurement/1-upload-uji.docx")
    doc_id = create_document(modul="procurement", nama="SOP Upload Uji", kode="1",
                              file_path="procurement/1-upload-uji.docx")

    new_doc = docx.Document()
    new_doc.add_heading("SOP Upload Uji", level=0)
    new_doc.add_heading("1. Tujuan", level=1)
    new_doc.add_paragraph("⚠ [PERLU VERIFIKASI: nominal baru]")
    buf = BytesIO()
    new_doc.save(buf)
    buf.seek(0)

    resp = client.post(
        f"/sop/{doc_id}/upload",
        data={"file": (buf, "1-upload-uji.docx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 302

    from modules.sop.service import get_document
    updated = get_document(doc_id)
    assert updated["perlu_verifikasi_count"] == 1
    assert updated["uploaded_by"] == "admin"

    conn_check = client.get(f"/sop/{doc_id}")
    assert b"nominal baru" in conn_check.data


def test_upload_rejects_non_docx_file(client):
    login(client)
    _write_sample_docx("procurement/1-reject-uji.docx")
    doc_id = create_document(modul="procurement", nama="SOP Reject Uji", kode="1",
                              file_path="procurement/1-reject-uji.docx")
    resp = client.post(
        f"/sop/{doc_id}/upload",
        data={"file": (BytesIO(b"not a docx"), "file.txt")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 302
    from modules.sop.service import get_document
    unchanged = get_document(doc_id)
    assert unchanged["file_path"] == "procurement/1-reject-uji.docx"
