import io
import re

from modules.management_report import routes as mr_routes
from modules.management_report import service as svc
from modules.management_report.parser import ParseError, ParseResult
from tests.test_mr_service import mkrow


def login(client, company_id=2):
    client.post("/auth/login", json={"username": "admin", "password": "Admin@123"})
    client.post("/select-company", data={"company_id": str(company_id)})


def test_page_requires_login(client):
    resp = client.get("/management-report/")
    assert resp.status_code == 302 and "/auth/login" in resp.headers["Location"]


def test_page_empty_state(client):
    login(client)
    resp = client.get("/management-report/")
    assert resp.status_code == 200
    assert b"Belum ada data" in resp.data


def _seed_bgt_screen():
    """AAA: Rental (months 1 and 9) + Gift (month 1); Period meta is Aug 2026."""
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", mm=1, mm1=1, bgt=10.0, act=4.0,
              category="01 - Rental, Tax", description="Rent"),
        mkrow(cc="AAA", coa="70105120", mm=9, mm1=9, bgt=6.0,
              category="01 - Rental, Tax", description="Rent"),
        mkrow(cc="AAA", coa="70106200", mm=1, mm1=1, bgt=3.0,
              category="06 - Gift", description="Gift A"),
        mkrow(cc="BBB", coa="70105120", mm=1, mm1=1, bgt=99.0,
              category="09 - Other", description="b-only"),
    ], meta={"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"})


def _screen_table(client, qs):
    return client.get("/management-report/table?" + qs).data.decode()


def test_page_shows_report_for_selected_cc(client):
    login(client)
    _seed_bgt_screen()
    resp = client.get("/management-report/?cc=AAA")
    html = resp.data.decode()
    assert resp.status_code == 200
    for text in ("Budget Full Year", "Budget YTD", "PTA YTD", "Actual YTD", "OPEX", "Rental, Tax", "Gift"):
        assert text in html
    assert "01 - Rental" not in _screen_table(client, "cc=AAA")   # the NN - prefix is cleaned
    assert "Grand Total" not in _screen_table(client, "cc=AAA")                    # sheet 1 has no grand total
    assert "b-only" not in html                         # other CC not shown
    assert "In Rp Mio" in html and "Aug 2026" in html and "YTD" in html
    assert "Buka semua" in html and "Tutup semua" in html


def test_screen_full_year_vs_ytd_and_hidden_details(client):
    login(client)
    _seed_bgt_screen()
    html = _screen_table(client, "cc=AAA&coa=70105120")
    assert "16,0" in html and "10,0" in html           # full year (months 1+9) vs YTD to Aug
    assert "- Bgt 2026 :" in html and "Rent" in html
    import re
    details = re.findall(r"<tr[^>]*mr-detail[^>]*>", html)
    assert details and all(" hidden" in d for d in details)
    assert 'class="mr-lvl-row' in html and 'class="mr-cat-row' in html
    # mm is the Period: the full year stays, YTD changes
    html = _screen_table(client, "cc=AAA&coa=70105120&mm=10")
    lvl = re.search(r'<tr class="mr-lvl-row[^>]*>.*?</tr>', html, re.S).group(0)   # detail rows keep their own 10,0
    assert "16,0" in lvl and "10,0" not in lvl and "Oct 2026" in html


def test_screen_filters_by_category_and_coa(client):
    login(client)
    _seed_bgt_screen()
    html = _screen_table(client, "cc=AAA&cat=06 - Gift")
    assert "Gift A" in html and "Rental, Tax" not in html
    html = _screen_table(client, "cc=AAA&coa=70106200")
    assert "Gift A" in html and "Rental, Tax" not in html


def test_screen_coa_subset_with_data_only_in_other_months_still_shows_rows(client):
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="111", mm=1, mm1=1, bgt=1.0, category="01 - Rental", description="early"),
        mkrow(cc="AAA", coa="222", mm=11, mm1=11, bgt=2.0, category="01 - Rental", description="late"),
    ], meta={"period": "2026-08-31"})
    html = _screen_table(client, "cc=AAA&coa=222")
    assert "Tidak ada data" not in html and "late" in html and "2,0" in html


def test_screen_empty_state_for_unknown_year(client):
    login(client)
    _seed_bgt_screen()
    assert "Tidak ada data untuk filter ini." in _screen_table(client, "cc=AAA&yy=1999")
    assert "Tidak ada data untuk filter ini." in client.get("/management-report/?cc=AAA&yy=1999").data.decode()


def test_screen_numbers_equal_sheet_1(client):
    from openpyxl import load_workbook
    from modules.management_report.formatting import id_number, id_percent
    login(client)
    _seed_bgt_screen()
    for qs in ("cc=AAA", "cc=AAA&mm=10", "cc=ALL", "cc=AAA&coa=70105120"):
        ws = load_workbook(io.BytesIO(client.get("/management-report/export?" + qs).data))["BGT Report per CC"]
        html = _screen_table(client, qs)
        row = re.search(r'<tr class="mr-lvl-row[^>]*>.*?</tr>', html, re.S).group(0)
        cells = [re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S)]
        expected = [id_number(ws[f"{c}10"].value) for c in "DEFGI"] + [id_number(ws["J10"].value), id_percent(ws["K10"].value),
                    id_number(ws["M10"].value), id_percent(ws["N10"].value)]
        assert cells[0] == "OPEX" and cells[1:] == expected, qs


def test_page_defaults_to_first_cc(client):
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [mkrow(cc="ZZZ"), mkrow(cc="AAA", description="first-cc-row")])
    assert b"first-cc-row" in client.get("/management-report/").data


def _sync(client, monkeypatch, rows, dropped=0):
    monkeypatch.setattr(mr_routes, "read_append1", lambda f: ParseResult(rows, dropped))
    return client.post("/management-report/sync",
                       data={"file": (io.BytesIO(b"x"), "Master.xlsb")},
                       content_type="multipart/form-data")


def test_sync_success(client, monkeypatch):
    login(client)
    resp = _sync(client, monkeypatch, [mkrow(), mkrow(cc="B")], dropped=3)
    body = resp.get_json()
    assert resp.status_code == 200 and body["ok"] is True and body["changed"] is True
    assert body["rows"] == 2 and body["dropped"] == 3
    assert body["summary"]["first"] is True
    assert svc.get_active_upload_id(2) == body["upload_id"]


def test_sync_again_without_changes_reports_no_change(client, monkeypatch):
    login(client)
    _sync(client, monkeypatch, [mkrow(), mkrow(cc="B")])
    resp = _sync(client, monkeypatch, [mkrow(cc="B"), mkrow()])
    body = resp.get_json()
    assert resp.status_code == 200 and body["ok"] is True and body["changed"] is False
    assert "tidak ada perubahan" in body["pesan"].lower()
    assert len(svc.list_uploads(2)) == 1


def test_sync_with_changes_reports_summary(client, monkeypatch):
    login(client)
    _sync(client, monkeypatch, [mkrow(cc="B", bgt=1.0)])
    body = _sync(client, monkeypatch, [mkrow(cc="B", bgt=4.0)]).get_json()
    assert body["changed"] is True and body["summary"]["lines"]["changed"] == 1
    assert body["summary"]["delta"]["bgt"] == 3.0
    assert "1 → 1" in body["pesan"]


def test_page_shows_sync_ui_and_last_summary(client, monkeypatch):
    login(client)
    html = client.get("/management-report/").data.decode()
    assert "Sync dari Excel" in html and "Sinkronisasi Excel" in html
    _sync(client, monkeypatch, [mkrow(cc="B", bgt=1.0)])
    _sync(client, monkeypatch, [mkrow(cc="B", bgt=4.0), mkrow(cc="C", bgt=2.0)])
    html = client.get("/management-report/").data.decode()
    assert "Perubahan sync terakhir" in html
    assert "CC terdampak (2)" in html


def test_sync_parse_error_and_missing_file(client, monkeypatch):
    login(client)

    def boom(f):
        raise ParseError("Kolom tidak ditemukan di sheet Append1: BGT")
    monkeypatch.setattr(mr_routes, "read_append1", boom)
    resp = client.post("/management-report/sync",
                       data={"file": (io.BytesIO(b"x"), "Master.xlsb")},
                       content_type="multipart/form-data")
    assert resp.status_code == 400 and "BGT" in resp.get_json()["pesan"]
    assert svc.get_active_upload_id(2) is None
    resp = client.post("/management-report/sync", data={}, content_type="multipart/form-data")
    assert resp.status_code == 400


def test_sync_rejects_non_xlsb_name(client):
    login(client)
    resp = client.post("/management-report/sync",
                       data={"file": (io.BytesIO(b"x"), "Master.xlsx")},
                       content_type="multipart/form-data")
    assert resp.status_code == 400 and ".xlsb" in resp.get_json()["pesan"]


def test_sync_requires_login(client):
    resp = client.post("/management-report/sync", data={}, content_type="multipart/form-data")
    assert resp.status_code == 401


def test_activate_endpoint(client):
    login(client)
    a = svc.save_upload(2, "a.xlsb", "u", [mkrow()])
    svc.save_upload(2, "b.xlsb", "u", [mkrow()])
    resp = client.post(f"/management-report/uploads/{a}/activate")
    assert resp.status_code == 200 and resp.get_json()["ok"] is True
    assert svc.get_active_upload_id(2) == a
    assert client.post("/management-report/uploads/9999/activate").status_code == 404

def test_non_releaser_role_denied_write_but_can_view(client):
    import bcrypt
    import database
    pw_hash = bcrypt.hashpw(b"Req@12345", bcrypt.gensalt(4)).decode()
    conn = database.get_conn()
    conn.execute(
        "INSERT INTO users (username, password_hash, role, must_change_pw) "
        "VALUES (?, ?, 'requester', 0)", ("req1", pw_hash))
    conn.commit()
    conn.close()
    resp = client.post("/auth/login", json={"username": "req1", "password": "Req@12345"})
    assert resp.status_code == 200
    client.post("/select-company", data={"company_id": "2"})

    assert client.get("/management-report/").status_code == 200
    up = client.post("/management-report/sync",
                     data={"file": (io.BytesIO(b"x"), "Master.xlsb")},
                     content_type="multipart/form-data")
    assert up.status_code == 403
    assert client.post("/management-report/uploads/1/activate").status_code == 403


# ---------- filters: All Cost Center, live table fragment, COA checklist ----------

def _seed_two_ccs():
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", category="01 - Rental", description="a-row", bgt=10.0, mm=1, coa="111"),
        mkrow(cc="BBB", category="01 - Rental", description="b-row", bgt=5.0, mm=2, coa="222"),
    ])


def test_all_cost_center_option_gives_consolidated_report(client):
    login(client)
    _seed_two_ccs()
    html = client.get("/management-report/?cc=ALL").data.decode()
    assert 'value="ALL" selected' in html
    assert "All Cost Center" in html
    assert "a-row" in html and "b-row" in html
    assert "15,0" in html                                   # grand total of both CCs


def test_table_fragment_follows_filters(client):
    login(client)
    _seed_two_ccs()
    resp = client.get("/management-report/table?cc=AAA")
    html = resp.data.decode()
    assert resp.status_code == 200 and "<html" not in html
    assert "a-row" in html and "b-row" not in html
    html = client.get("/management-report/table?cc=ALL&mm=2").data.decode()
    assert "b-row" in html and "a-row" in html          # mm is the Period (year to date): month 1 counts too
    html = client.get("/management-report/table?cc=ALL&coa=111").data.decode()
    assert "a-row" in html and "b-row" not in html
    assert "a-row" in client.get("/management-report/table?cc=AAA&mm=2").data.decode()
    assert "Tidak ada data" in client.get("/management-report/table?cc=AAA&yy=1999").data.decode()


def test_table_fragment_requires_login(client):
    assert client.get("/management-report/table?cc=ALL").status_code == 302


def test_coa_filter_is_a_checklist_with_select_all(client):
    login(client)
    _seed_two_ccs()
    html = client.get("/management-report/?cc=ALL").data.decode()
    assert html.count('class="mr-coa-item"') == 2 and 'id="mr-coa-all"' in html
    assert 'value="111" checked' in html and 'value="222" checked' in html      # none filtered: all ticked
    html = client.get("/management-report/?cc=ALL&coa=111").data.decode()
    assert 'value="111" checked' in html and 'value="222" checked' not in html


def test_export_all_cost_centers_in_one_workbook(client):
    from openpyxl import load_workbook
    login(client)
    _seed_two_ccs()
    resp = client.get("/management-report/export?cc=ALL")
    assert resp.status_code == 200 and "ALL.xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.worksheets[0]["C1"].value.endswith("All")
    append_values = [c.value for row in wb["Append"].iter_rows(min_row=2) for c in row]
    assert "b-row" in append_values and "a-row" in append_values


def test_sync_stores_workbook_metadata(client, monkeypatch):
    login(client)
    monkeypatch.setattr(mr_routes, "read_meta",
                        lambda f: {"cc_names": {"B": "Beta"}, "period": "2026-08-31"})
    _sync(client, monkeypatch, [mkrow(cc="B")])
    assert svc.get_active_meta(2) == {"cc_names": {"B": "Beta"}, "period": "2026-08-31"}


HO_META = {"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"}


def _seed_ho():
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", bgt=3.2, act=0.5, mm=1),
        mkrow(cc="BBB", coa="71000110", act=-97.5, mm=1),
    ], meta=HO_META)


def test_page_has_three_tabs_and_the_report_title_uses_the_cc_description(client):
    login(client)
    _seed_ho()
    html = client.get("/management-report/?cc=AAA").data.decode()
    for tab in ("tab-sync", "tab-report", "tab-ho"):
        assert f'data-tab="{tab}"' in html and f'id="{tab}"' in html
    assert "President Office &amp; ETF – Alpha Unit" in html
    assert "President Office &amp; ETF – All" in client.get("/management-report/?cc=ALL").data.decode()


def test_ho_fragment_shows_template_rows_and_amounts(client):
    login(client)
    _seed_ho()
    resp = client.get("/management-report/ho-table?cc=AAA")
    html = resp.data.decode()
    assert resp.status_code == 200 and "<html" not in html
    assert "Equipment Rent Expense - 3rd Party" in html and "3.200" in html
    assert "270-01-010" in html                      # GL code shown like the Excel format 000-00-000
    assert "s/d Aug" in html and "Alpha Unit" in html
    assert "Total Capex" in html and "Grand Total included Depreciation" in html


def test_ho_fragment_warns_about_coa_outside_the_template(client):
    login(client)
    _seed_ho()
    html = client.get("/management-report/ho-table?cc=BBB").data.decode()
    assert "COA di luar template HO" in html and "71000110" in html and "(97.500)" in html
    assert "COA di luar template HO" not in client.get("/management-report/ho-table?cc=AAA").data.decode()


def test_ho_fragment_login_and_empty_state(client):
    assert client.get("/management-report/ho-table?cc=ALL").status_code == 302
    login(client)
    assert "Tidak ada data" in client.get("/management-report/ho-table?cc=ALL").data.decode()


def test_ho_fragment_follows_month_filter_and_coa(client):
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [mkrow(cc="AAA", coa="70105120", bgt=3.2, mm=1),
                                       mkrow(cc="AAA", coa="70106200", bgt=1.0, mm=2)])
    html = client.get("/management-report/ho-table?cc=AAA&mm=2").data.decode()
    assert "Local Travel Expense" in html and "1.000" in html     # mm=2 row is counted
    html = client.get("/management-report/ho-table?cc=AAA&mm=12").data.decode()
    assert "Tidak ada data" in html


# ---------- look & feel shared with the Sponsorship report pages ----------

def test_report_and_ho_tabs_are_full_bleed_and_the_filter_bar_is_collapsible(client):
    login(client)
    _seed_ho()
    html = client.get("/management-report/?cc=AAA").data.decode()
    assert 'id="tab-report" data-full-bleed="true"' in html
    assert 'id="tab-ho" data-full-bleed="true"' in html
    assert 'id="tab-sync" data-full-bleed' not in html            # the sync tab is a normal page
    assert 'class="budget-toolbar filter-bar collapsible"' in html
    assert 'id="mr-filter-panel"' in html                         # stable id: the app persists open/closed state by it
    assert "budget.css" in html and "sponsorship.css" in html


def test_report_table_uses_the_sponsorship_table_classes(client):
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", category="06 - Gift", bgt=33.5, act=7.0, description="Gift A")])
    html = client.get("/management-report/table?cc=AAA").data.decode()
    assert 'class="budget-table' in html
    assert 'class="mr-lvl-row' in html and 'class="mr-cat-row' in html
    assert "33,5" in html                                          # Indonesian number style


def test_ho_table_uses_the_sponsorship_table_classes(client):
    login(client)
    _seed_ho()
    html = client.get("/management-report/ho-table?cc=AAA").data.decode()
    assert html.count('class="budget-table') == 3                  # capex, opex, summary
    assert 'class="sp-report-section-band"' in html                # group headers
    assert 'class="sp-report-subtotal-row"' in html and 'class="sp-report-grandtotal-row"' in html
    assert "3.200" in html and "(97.500)" not in html


def test_coa_filter_uses_the_shared_dropdown_check_markup(client):
    login(client)
    _seed_ho()
    html = client.get("/management-report/?cc=AAA").data.decode()
    assert 'class="budget-dropdown-check"' in html and 'class="budget-dropdown-btn"' in html
    assert 'class="budget-dropdown-content"' in html


# ---------- Data Append (Data Sync tab) ----------

def _seed_append():
    rows = [mkrow(cc="AAA", coa="70105120", description=f"row {i}", mm=1, bgt=1.5)
            for i in range(150)]
    rows.append(mkrow(cc="BBB", coa="70106200", description="travel 100% <b>x</b>", mm=2, bgt=2.0))
    svc.save_upload(2, "a.xlsb", "u", rows)


def test_data_sync_tab_has_the_append_section(client):
    login(client)
    _seed_append()
    html = client.get("/management-report/?tab=sync").data.decode()
    assert 'id="mr-append"' in html and 'id="mr-append-cc"' in html and 'id="mr-append-q"' in html
    assert "Export Append (Excel)" in html


def test_append_fragment_pages_filters_and_escapes(client):
    login(client)
    _seed_append()
    get = lambda qs: client.get("/management-report/append-table" + qs).data.decode()
    html = get("")
    assert "<html" not in html and "151 baris" in html and "Halaman 1 dari 2" in html
    assert "row 0<" in html and "row 99<" in html and "row 100<" not in html
    page2 = get("?page=2")
    assert "Halaman 2 dari 2" in page2 and "row 149<" in page2 and "row 0<" not in page2
    assert "Halaman 2 dari 2" in get("?page=99")                     # clamped
    assert "Halaman 1 dari 2" in get("?page=abc") and "Halaman 1 dari 2" in get("?page=0")
    by_cc = get("?cc=BBB")
    assert ">1 baris<" in by_cc and "travel 100%" in by_cc
    assert "151 baris" in get("?cc=NOPE")                            # unknown CC = all
    assert "151 baris" in get("?cc=ALL")
    assert ">1 baris<" in get("?q=travel") and ">1 baris<" in get("?q=%25")   # '%' matches literally
    assert "&lt;b&gt;x&lt;/b&gt;" in by_cc and "<b>x</b>" not in by_cc    # escaped
    assert "Tidak ada baris" in get("?q=zzz")


def test_append_fragment_empty_state_and_login(client):
    assert client.get("/management-report/append-table").status_code == 302
    login(client)
    assert "Belum ada data" in client.get("/management-report/append-table").data.decode()


def test_append_export_downloads_the_filtered_rows(client):
    from openpyxl import load_workbook
    login(client)
    _seed_append()
    resp = client.get("/management-report/append/export")
    assert resp.status_code == 200 and "Append.xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Append"] and wb["Append"].max_row == 1 + 151
    only_b = load_workbook(io.BytesIO(client.get("/management-report/append/export?cc=BBB").data))
    assert only_b["Append"].max_row == 2
    assert load_workbook(io.BytesIO(
        client.get("/management-report/append/export?q=travel").data))["Append"].max_row == 2
    assert client.get("/management-report/append/export?q=zzz").status_code == 404


def test_append_export_login_and_empty_snapshot(client):
    assert client.get("/management-report/append/export").status_code == 302
    login(client)
    assert client.get("/management-report/append/export").status_code == 404


def test_report_export_has_an_append_sheet_that_follows_the_report_filters(client):
    from openpyxl import load_workbook
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="111", mm=1, bgt=1.0), mkrow(cc="AAA", coa="222", mm=2, bgt=2.0),
        mkrow(cc="BBB", coa="111", mm=2, bgt=4.0)])
    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA").data))
    assert wb.sheetnames == ["BGT Report per CC", "BGT HO Report per CC", "Append"]
    assert wb["Append"].max_row == 1 + 2
    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=ALL&mm=2&coa=111").data))
    ws = wb["Append"]
    assert ws.max_row == 1 + 1 and ws.cell(2, 4).value == "BBB"          # CC, MM and COA filters apply
    all_ccs = load_workbook(io.BytesIO(client.get("/management-report/export?cc=ALL").data))["Append"]
    assert all_ccs.max_row == 1 + 3


def test_append_fragment_survives_extreme_pages_and_matches_a_backslash_literally(client):
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [mkrow(cc="AAA", description="path a\\b end"),
                                       mkrow(cc="AAA", description="plain")])
    get = lambda qs: client.get("/management-report/append-table" + qs)
    resp = get("?page=99999999999999999999999")                 # far beyond SQLite's integer range
    assert resp.status_code == 200 and "Halaman 1 dari 1" in resp.data.decode()
    assert get("?page=-5").status_code == 200
    assert ">1 baris<" in get("?q=a%5Cb").data.decode()          # backslash is a literal character
    assert "Tidak ada baris" in get("?q=%5C%5C").data.decode()   # a double backslash matches nothing


def test_export_first_sheet_follows_the_reference_layout(client):
    from openpyxl import load_workbook
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", mm=1, mm1=1, bgt=10.0, act=4.0,
              category="01 - Rental, Tax", description="Rent"),
        mkrow(cc="AAA", coa="70105120", mm=9, mm1=9, bgt=6.0,
              category="01 - Rental, Tax", description="Rent"),
    ], meta={"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"})
    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA").data))
    assert wb.sheetnames == ["BGT Report per CC", "BGT HO Report per CC", "Append"]
    ws = wb["BGT Report per CC"]
    assert ws["C1"].value == "President Office & ETF – Alpha Unit"
    assert ws["A2"].value == "AAA" and ws["A3"].value == "ETF - Alpha Unit"
    assert ws["C10"].value == "OPEX"
    assert (ws["D10"].value, ws["E10"].value, ws["I10"].value) == (16, 10, 4)
    # mm filter sets the Period, full year stays the whole year
    ws = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA&mm=10").data))["BGT Report per CC"]
    assert (ws["D10"].value, ws["E10"].value) == (16, 16)
    assert ws["D4"].value.month == 10


def test_export_with_out_of_range_month_does_not_crash(client):
    from openpyxl import load_workbook
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [mkrow(cc="AAA", coa="111", mm=1, bgt=1.0)],
                    meta={"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"})
    resp = client.get("/management-report/export?cc=AAA&mm=13")
    assert resp.status_code == 200
    assert load_workbook(io.BytesIO(resp.data)).sheetnames[0] == "BGT Report per CC"


def test_export_all_cc_header_values(client):
    from openpyxl import load_workbook
    login(client)
    _seed_two_ccs()
    ws = load_workbook(io.BytesIO(client.get("/management-report/export?cc=ALL").data))["BGT Report per CC"]
    assert ws["A2"].value == "All"
    assert ws["A3"].value.endswith("- All")


def test_export_all_zip_bgt_sheet_content(client):
    import zipfile
    from openpyxl import load_workbook
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", mm=1, mm1=1, bgt=10.0, act=4.0,
              category="01 - Rental, Tax", description="Rent"),
    ], meta={"cc_names": {"AAA": "Alpha Unit"}, "period": "2026-08-31"})
    resp = client.get("/management-report/export-all")
    assert resp.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(resp.data))
    assert zf.namelist() == ["AAA.xlsx"]
    ws = load_workbook(io.BytesIO(zf.read("AAA.xlsx")))["BGT Report per CC"]
    assert ws["C1"].value == "President Office & ETF – Alpha Unit"
    assert ws["A2"].value == "AAA" and ws["A3"].value == "ETF - Alpha Unit"
    assert ws["C10"].value == "OPEX"


def test_export_all_includes_ccs_with_lines_only_in_other_months(client):
    import zipfile
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="111", mm=1, bgt=1.0),
        mkrow(cc="BBB", coa="111", mm=2, bgt=2.0),
        mkrow(cc="CCC", coa="111", mm=12, bgt=3.0),
    ])
    resp = client.get("/management-report/export-all?mm=5")        # no CC has data in month 5
    assert resp.status_code == 200
    # AAA and BBB have lines up to May (year to date); CCC only in December
    assert sorted(zipfile.ZipFile(io.BytesIO(resp.data)).namelist()) == ["AAA.xlsx", "BBB.xlsx"]


def test_company_code_lookup():
    assert mr_routes._company_code(1) == "SMT"
    assert mr_routes._company_code(2) == "ETF"
    assert mr_routes._company_code(99) == ""


def test_ho_follows_the_category_and_coa_filters_on_screen_and_in_the_export(client):
    import zipfile
    from openpyxl import load_workbook
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", category="01 - Rental", bgt=3.2, description="Rent"),
        mkrow(cc="AAA", coa="70102200", category="Salary", bgt=9.0, description="Pay")],
        meta={"cc_names": {"AAA": "Alpha"}, "period": "2026-08-31"})

    def ho_numbers(wb):
        ws = wb["BGT HO Report per CC"]
        return {ws.cell(r, 1).value: ws.cell(r, 4).value for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value in (70105120, 70102200)}

    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA").data))
    assert ho_numbers(wb) == {70105120: 3200.0, 70102200: 9000.0}
    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA&cat=01 - Rental").data))
    assert ho_numbers(wb) == {70105120: 3200.0, 70102200: 0}                  # Salary unticked
    wb = load_workbook(io.BytesIO(client.get("/management-report/export?cc=AAA&coa=70102200").data))
    assert ho_numbers(wb) == {70105120: 0, 70102200: 9000.0}
    zf = zipfile.ZipFile(io.BytesIO(client.get("/management-report/export-all?cat=01 - Rental").data))
    assert ho_numbers(load_workbook(io.BytesIO(zf.read("AAA.xlsx")))) == {70105120: 3200.0, 70102200: 0}
    # screen
    html = client.get("/management-report/ho-table?cc=AAA&cat=01 - Rental").data.decode()
    assert "Equipment Rent Expense" in html and "3.200" in html and "9.000" not in html
    html = client.get("/management-report/ho-table?cc=AAA&coa=70102200").data.decode()
    assert "9.000" in html and "3.200" not in html


def test_var_columns_are_green_under_budget_and_red_over_budget(client):
    import re
    login(client)
    svc.save_upload(2, "a.xlsb", "u", [
        mkrow(cc="AAA", coa="70105120", category="01 - Rental", bgt=10.0, act=4.0, description="Under"),
        mkrow(cc="AAA", coa="70106100", category="02 - Transport", bgt=2.0, act=9.0, description="Over")])
    html = client.get("/management-report/table?cc=AAA").data.decode()

    def classes_of(label):
        row = re.search(r"<tr[^>]*>(?:(?!</tr>).)*?" + label + r"(?:(?!</tr>).)*</tr>", html, re.S).group(0)
        return [c.strip() for c in re.findall(r'<td class="num-right([^"]*)"', row)]

    rental, transport = classes_of("Rental"), classes_of("Transport")
    assert rental[5] == "mr-under" and rental[6] == "mr-under"          # Var amount and Var %
    assert transport[5] == "mr-over" and transport[6] == "mr-over"
