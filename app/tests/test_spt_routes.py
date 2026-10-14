# tests/test_spt_routes.py
import io
import re

import bcrypt
import openpyxl

import database
from modules.management_report import category_map as cm
from modules.spt_badan import service as sv
from tests.test_mrc_routes import login
from tests.test_spt_service import SPT_GAJI, new_period, seed_gl, seed_mapping

BASE = "/spt-badan"


def cells(html, table_id):
    block = html.split(f'id="{table_id}"', 1)[1].split("</table>", 1)[0]
    return [[re.sub(r"<.*?>", "", c).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", r, re.S)]
            for r in re.findall(r"<tr[^>]*>(.*?)</tr>", block, re.S) if "<td" in r]


def make_requester(client):
    pw = bcrypt.hashpw(b"Req@12345", bcrypt.gensalt(4)).decode()
    conn = database.get_conn()
    conn.execute("INSERT INTO users (username, password_hash, role, must_change_pw) "
                 "VALUES (?, ?, 'requester', 0)", ("req1", pw))
    conn.commit()
    conn.close()
    client.post("/auth/login", json={"username": "req1", "password": "Req@12345"})
    client.post("/select-company", data={"company_id": "2"})


def ready(client):
    login(client)
    seed_mapping(); seed_gl()
    pid = new_period()
    sid = client.post(f"{BASE}/{pid}/snapshot").get_json()["snapshot_id"]
    return pid, sid


def test_pages_require_login(client):
    for url in (BASE + "/", BASE + "/1", BASE + "/1/laba-rugi", BASE + "/1/pengecualian", BASE + "/1/kontrol"):
        r = client.get(url)
        assert r.status_code == 302 and "/auth/login" in r.headers["Location"], url


def test_index_lists_periods_and_admin_can_create(client):
    login(client)
    html = client.get(BASE + "/").data.decode()
    assert "SPT Badan" in html and 'id="spt-create-btn"' in html
    r = client.post(BASE + "/periods", json={"entity": "smt", "tax_year": 2026})
    assert r.status_code == 200 and r.get_json()["ok"]
    html = client.get(BASE + "/").data.decode()
    assert cells(html, "spt-periods")[0][:2] == ["SMT", "2026"] and "Simulasi YTD" in html
    dup = client.post(BASE + "/periods", json={"entity": "smt", "tax_year": 2026})
    assert dup.status_code == 400 and "sudah ada" in dup.get_json()["pesan"]
    assert client.post(BASE + "/periods", json={"entity": "smt", "tax_year": "x"}).status_code == 400


def test_non_admin_can_view_but_not_create_or_snapshot(client):
    seed_mapping(); seed_gl()
    pid = new_period()
    make_requester(client)
    html = client.get(BASE + "/").data.decode()
    assert 'id="spt-create-btn"' not in html and client.get(f"{BASE}/{pid}").status_code == 200
    assert client.post(BASE + "/periods", json={"entity": "smt", "tax_year": 2027}).status_code == 403
    assert client.post(f"{BASE}/{pid}/snapshot").status_code == 403
    assert sv.get_current_snapshot(pid) is None and len(sv.list_periods()) == 1


def test_unknown_period_is_404(client):
    login(client)
    for url in ("/999", "/999/laba-rugi", "/999/pengecualian", "/999/kontrol", "/999/baris",
                "/999/pengecualian.xlsx"):
        assert client.get(BASE + url).status_code == 404, url
    assert client.post(BASE + "/999/snapshot").status_code == 404


def test_detail_before_and_after_snapshot(client):
    login(client)
    seed_mapping(); seed_gl()
    pid = new_period()
    html = client.get(f"{BASE}/{pid}").data.decode()
    assert "Belum ada snapshot" in html and 'id="spt-snapshot-btn"' in html and "Simulasi YTD" in html
    j = client.post(f"{BASE}/{pid}/snapshot").get_json()
    assert j["ok"] and j["blockers"] == 0 and "8 baris" in j["pesan"]
    html = client.get(f"{BASE}/{pid}").data.decode()
    assert "2026-01-02" in html and "Laba sebelum pajak" in html and "Belum ada snapshot" not in html
    assert [r[1] for r in cells(html, "spt-steps")] == ["lulus", "peringatan", "lulus", "peringatan"]


def test_snapshot_without_gl_returns_a_clear_error(client):
    login(client)
    seed_mapping()
    pid = new_period()
    r = client.post(f"{BASE}/{pid}/snapshot")
    assert r.status_code == 400 and "Sync General Ledger" in r.get_json()["pesan"]


def test_snapshot_failure_returns_json_500(client, monkeypatch):
    login(client)
    seed_mapping(); seed_gl()
    pid = new_period()
    monkeypatch.setattr(sv, "take_snapshot", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("db locked")))
    r = client.post(f"{BASE}/{pid}/snapshot")
    assert r.status_code == 500 and "db locked" in r.get_json()["pesan"]


def test_income_statement_page_shows_numbers_and_drill_links(client):
    pid, sid = ready(client)
    html = client.get(f"{BASE}/{pid}/laba-rugi").data.decode()
    rows = {r[0]: r[1] for r in cells(html, "spt-statement")}
    assert rows["Pendapatan bersih"] == "1.000" and rows["Laba sebelum pajak penghasilan"] == "440"
    assert rows["Laba (rugi) bersih"] == "360" and "Simulasi YTD" in html
    assert f"/spt-badan/{pid}/baris?" in html


def test_drill_page_total_equals_the_statement_row(client):
    pid, sid = ready(client)
    html = client.get(f"{BASE}/{pid}/baris",
                      query_string={"cats": ["HARGA POKOK PENJUALAN", "Beban Operasional"]}).data.decode()
    assert [r[2] for r in cells(html, "spt-lines")] and "600" in html            # 300+100+200
    one = client.get(f"{BASE}/{pid}/baris", query_string={"cats": "HARGA POKOK PENJUALAN", "spt": "BIAYA SEWA"})
    assert [r[2] for r in cells(one.data.decode(), "spt-lines")] == ["5-2000"]
    assert client.get(f"{BASE}/{pid}/baris").status_code == 200                  # no category: empty list


def test_exceptions_page_groups_filters_and_exports(client):
    pid, sid = ready(client)
    html = client.get(f"{BASE}/{pid}/pengecualian").data.decode()
    groups = cells(html, "spt-exc-groups")
    assert {(g[0], g[2]) for g in groups} == {("pph_kosong", "6-1000"), ("pph_bukan_jenis", "8-2000")}
    only = client.get(f"{BASE}/{pid}/pengecualian?kind=pph_bukan_jenis").data.decode()
    assert [r[2] for r in cells(only, "spt-exc-rows")] == ["8-2000"]
    x = client.get(f"{BASE}/{pid}/pengecualian.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["Content-Type"]
    ws = openpyxl.load_workbook(io.BytesIO(x.data)).active
    assert [c.value for c in ws[1]][:4] == ["Jenis", "Tingkat", "Akun", "Tanggal"] and ws.max_row == 3


def test_controls_page_lists_c1_to_c9_and_marks_stale(client):
    pid, sid = ready(client)
    html = client.get(f"{BASE}/{pid}/kontrol").data.decode()
    assert [r[0] for r in cells(html, "spt-controls")] == ["C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8", "C9"]
    assert "GL berubah sejak snapshot" not in html
    seed_gl(hash_="hash-2")
    html = client.get(f"{BASE}/{pid}").data.decode()
    assert "GL berubah sejak snapshot" in html and "ambil snapshot ulang" in html.lower()
    assert "GL berubah sejak snapshot" in client.get(f"{BASE}/{pid}/kontrol").data.decode()


def test_module_tab_is_in_the_dashboard_navigation(client):
    login(client)
    html = client.get(BASE + "/").data.decode()
    assert "SPT Badan" in html and "/spt-badan/" in html
