import io
import openpyxl
from modules.sponsorship.service import seed_perusahaan, create_transaksi
from modules.sponsorship.exports import export_mr_excel

SMT = 1


def _seed():
    seed_perusahaan([("SMART", "AGRI", "AGRI")])


def _payload(**overrides):
    payload = {
        "cat": "Sponsorship", "description": "Sosial",
        "tanggal_form_sponsor": "2026-01-20", "cat1": "SMART",
        "lembaga_perorangan": "Yayasan Test", "perihal": "Bantuan Banjir",
        "amount_request": 25000000,
    }
    payload.update(overrides)
    return payload


def test_export_mr_excel_has_expected_headers(app):
    _seed()
    xls_bytes = export_mr_excel(SMT, ["2026"])
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == [
        "Account Management", "ID#", "Date", "Allocation Memo", "Budget Activity",
        "COST CENTER", "DEBIT - CREDIT", "Source", "COA HO", "COST CENTER2",
        "Month", "CC0", "CC1", "BC1", "New Code", "Memo", "BC",
    ]


def test_export_mr_excel_matches_worked_example_constants(app):
    _seed()
    create_transaksi(SMT, _payload(
        perihal="Partisipasi Pemasangan Lampion dan Lilin Tahun Baru Imlek 2026",
        tanggal_form_sponsor="2026-01-20", amount_request=25000000,
    ))
    create_transaksi(SMT, _payload(
        perihal="Prasasti Membership",
        tanggal_form_sponsor="2026-01-20", amount_request=500000000,
    ))
    xls_bytes = export_mr_excel(SMT, ["2026"])
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    row1 = [c.value for c in ws[2]]
    row2 = [c.value for c in ws[3]]
    assert row1[0] == 70107800          # Account Management
    assert row1[5] == "POSPON"          # COST CENTER
    assert row1[6] == 25000000          # DEBIT - CREDIT
    assert row1[7] == "Sponsor"         # Source
    assert row1[8] == 70107800          # COA HO
    assert row1[9] == "$POCCOM"         # COST CENTER2
    assert row1[10] == 1                # Month
    assert row1[11] == "26"             # CC0
    assert row1[12] == 5                # CC1
    assert row1[13] == "-25-"           # BC1
    assert row1[14] == "A26-5-25-1"     # New Code
    assert row1[16] == "Sponsorship - Non Building"  # BC
    assert row2[6] == 500000000         # DEBIT - CREDIT
    assert row2[14] == "A26-5-25-2"     # New Code increments


def test_export_mr_excel_empty_data_still_generates_file(app):
    _seed()
    xls_bytes = export_mr_excel(SMT, ["2026"])
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    assert ws["A1"].value == "Account Management"
    assert ws.max_row == 1
