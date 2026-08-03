path = r"C:\Financehub\app\modules\sponsorship\exports.py"
with open(path, "r", encoding="utf-8") as f:
    c = f.read()

import re

# We will just replace the entire file since it only contains these exports
# Wait, I added export_data_excel and export_outstanding_excel to exports.py!
# So I should only replace export_report_excel and the constants at the top.

# Let's extract everything from the start up to export_data_excel
pattern = re.compile(r"^(.*?)def export_data_excel", re.DOTALL)
match = pattern.search(c)

new_code = """import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from modules.sponsorship.service import get_sponsorship_report, PILLAR_COLUMNS

# --- Styles ---
ARIAL_11 = Font(name='Arial', size=11)
ARIAL_11_BOLD = Font(name='Arial', size=11, bold=True)
ARIAL_11_WHITE_BOLD = Font(name='Arial', size=11, bold=True, color="FFFFFF")

DARK_BLUE_FILL = PatternFill("solid", fgColor="11325B")  # Dark blue for headers
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")     # Yellow for categories/subtotals
APP_FILL = PatternFill("solid", fgColor="FCE4D6")        # Light peach/pink
AGRI_FILL = PatternFill("solid", fgColor="F4B183")       # Light orange
EF_FILL = PatternFill("solid", fgColor="A9D08E")         # Light green
PROP_FILL = PatternFill("solid", fgColor="B4A7D6")       # Light purple

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(wrap_text=True, vertical="top")

THIN_BORDER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

_NUM_FMT = "#,##0"
_COLUMN_LABELS = ["NO", "LEMBAGA / PERORANGAN", "PERIHAL", "APP", "AGRI", "ENERGY & FINANCE", "PROPERTY", "TOTAL"]

def export_report_excel(company_id: int, tahun: str, cat4_mode: str) -> bytes:
    data = get_sponsorship_report(company_id, tahun, cat4_mode)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    # Set default font for all cells (doesn't always work perfectly, so we'll set explicitly)
    
    # Title row
    ws.append([f"RANGKUMAN SPONSORSHIP - in Rp Full Amount"])
    ws["A1"].font = ARIAL_11_BOLD
    ws.append([])

    def _write_section_headers():
        # Dark Blue Header Row
        ws.append([f"RANGKUMAN SPONSORSHIP {data['tahun']} - in Rp Full Amount"])
        row = ws.max_row
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.fill = DARK_BLUE_FILL
            cell.font = ARIAL_11_WHITE_BOLD
            cell.border = THIN_BORDER

        # Column Labels Row
        ws.append(_COLUMN_LABELS)
        row = ws.max_row
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.font = ARIAL_11_BOLD
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            
            if col in [1, 2, 3, 8]:
                cell.fill = DARK_BLUE_FILL
                cell.font = ARIAL_11_WHITE_BOLD
            elif col == 4: cell.fill = APP_FILL
            elif col == 5: cell.fill = AGRI_FILL
            elif col == 6: cell.fill = EF_FILL
            elif col == 7: cell.fill = PROP_FILL

    if not data["sections"]:
        _write_section_headers()

    for section in data["sections"]:
        _write_section_headers()
        
        # Category Row (Yellow)
        ws.append([section["description"].upper()])
        row = ws.max_row
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.fill = YELLOW_FILL
            cell.font = ARIAL_11_BOLD
            cell.border = THIN_BORDER

        # Data Rows
        for r in section["rows"]:
            ws.append([
                r["no"], r["lembaga_perorangan"], r["perihal"] or "",
                r["app_amount"], r["agri_amount"], r["energy_finance_amount"], r["property_amount"], r["total"]
            ])
            row = ws.max_row
            for col in range(1, 9):
                cell = ws.cell(row=row, column=col)
                cell.font = ARIAL_11
                cell.border = THIN_BORDER
                if col >= 4:
                    cell.number_format = _NUM_FMT
                    cell.alignment = ALIGN_RIGHT
            
            ws.cell(row=row, column=1).alignment = ALIGN_CENTER
            ws.cell(row=row, column=2).alignment = ALIGN_WRAP
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP

        # Subtotal Row
        sub = section["subtotal"]
        ws.append([
            f"TOTAL {section['description'].upper()}", "", "",
            sub["app_amount"], sub["agri_amount"], sub["energy_finance_amount"], sub["property_amount"], sub["total"]
        ])
        row = ws.max_row
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.fill = YELLOW_FILL
            cell.font = ARIAL_11_BOLD
            cell.border = THIN_BORDER
            if col >= 4:
                cell.number_format = _NUM_FMT
                cell.alignment = ALIGN_RIGHT

    # Grand Total
    gt = data["grand_total"]
    ws.append([
        "TOTAL SPONSORSHIP (ACTUAL)", "", "",
        gt["app_amount"], gt["agri_amount"], gt["energy_finance_amount"], gt["property_amount"], gt["total"]
    ])
    row = ws.max_row
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        cell.fill = DARK_BLUE_FILL
        cell.font = ARIAL_11_WHITE_BOLD
        cell.border = THIN_BORDER
        if col >= 4:
            cell.number_format = _NUM_FMT
            cell.alignment = ALIGN_RIGHT

    # PCT Allocation
    pct = data["pct_allocation"]
    ws.append([
        "% Allocation", "", "",
        f"{pct['app_amount']}%", f"{pct['agri_amount']}%", f"{pct['energy_finance_amount']}%", f"{pct['property_amount']}%", ""
    ])
    row = ws.max_row
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        cell.font = ARIAL_11_BOLD
        cell.border = THIN_BORDER
        if col >= 4:
            cell.alignment = ALIGN_RIGHT

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 50
    for col_letter in ("D", "E", "F", "G", "H"):
        ws.column_dimensions[col_letter].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

"""

if match:
    # Everything after export_data_excel
    rest = c[match.end(1):]
    new_c = new_code + "\n" + rest
    with open(path, "w", encoding="utf-8") as f:
        f.write(new_c)
    print("done")
else:
    print("pattern not found")
