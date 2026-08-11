import docx
from modules.sop.renderer import parse_docx


def _build_test_docx(path):
    document = docx.Document()
    document.add_heading("SOP Contoh Uji", level=0)
    document.add_heading("1. Tujuan", level=1)
    document.add_paragraph("Menjadi panduan uji renderer.", style="List Bullet")
    document.add_paragraph("Poin kedua.", style="List Bullet")
    document.add_heading("4. Kebijakan", level=1)
    document.add_heading("4.1 Pihak Terkait", level=2)
    p = document.add_paragraph()
    p.add_run("Karyawan").bold = True
    document.add_paragraph("Langkah pertama.", style="List Number")
    document.add_paragraph("Langkah kedua.", style="List Number")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Versi", "Tanggal", "Deskripsi Perubahan", "Oleh"
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text, row[3].text = "1.0", "1 Jan 2026", "Awal", "HR"
    document.save(path)


def test_parse_docx_extracts_title(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    assert result["title"] == "SOP Contoh Uji"


def test_parse_docx_extracts_section_headings_in_order(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    headings = [s["heading"] for s in result["sections"]]
    assert headings == ["1. Tujuan", "4. Kebijakan", "4.1 Pihak Terkait"]


def test_parse_docx_renders_bullet_list_as_ul(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    tujuan = next(s for s in result["sections"] if s["heading"] == "1. Tujuan")
    assert "<ul><li>Menjadi panduan uji renderer.</li><li>Poin kedua.</li></ul>" == tujuan["html"]


def test_parse_docx_renders_numbered_list_as_ol(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    pihak = next(s for s in result["sections"] if s["heading"] == "4.1 Pihak Terkait")
    assert "<ol><li>Langkah pertama.</li><li>Langkah kedua.</li></ol>" in pihak["html"]


def test_parse_docx_renders_bold_run_as_strong(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    pihak = next(s for s in result["sections"] if s["heading"] == "4.1 Pihak Terkait")
    assert "<strong>Karyawan</strong>" in pihak["html"]


def test_parse_docx_extracts_revisi_rows_separately(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    assert result["revisi_rows"] == [
        {"versi": "1.0", "tanggal": "1 Jan 2026", "deskripsi": "Awal", "oleh": "HR"}
    ]
    for section in result["sections"]:
        assert "Versi" not in section["html"]


def test_parse_docx_captures_pre_heading_content_as_pendahuluan(tmp_path):
    path = str(tmp_path / "preamble.docx")
    document = docx.Document()
    document.add_heading("SOP Dengan Preamble", level=0)
    document.add_paragraph("Catatan pembuka sebelum bab pertama.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Kolom A", "Kolom B"
    row = table.add_row().cells
    row[0].text, row[1].text = "Nilai A", "Nilai B"
    document.add_heading("1. Tujuan", level=1)
    document.add_paragraph("Isi tujuan.")
    document.save(path)

    result = parse_docx(path)
    assert result["sections"][0]["heading"] == "Pendahuluan"
    assert result["sections"][0]["level"] == 1
    assert "Catatan pembuka sebelum bab pertama." in result["sections"][0]["html"]
    assert "<table" in result["sections"][0]["html"]
    assert "Nilai A" in result["sections"][0]["html"]
    assert result["sections"][1]["heading"] == "1. Tujuan"


def test_parse_docx_no_preamble_content_does_not_add_phantom_section(tmp_path):
    path = str(tmp_path / "contoh.docx")
    _build_test_docx(path)
    result = parse_docx(path)
    assert result["sections"][0]["heading"] != "Pendahuluan"


def test_parse_docx_table_with_partial_versi_header_is_not_revisi(tmp_path):
    path = str(tmp_path / "partial-versi.docx")
    document = docx.Document()
    document.add_heading("SOP Tabel Versi Parsial", level=0)
    document.add_heading("1. Tujuan", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Versi", "Nama", "Catatan"
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = "1.0", "Budi", "Tidak relevan"
    document.save(path)

    result = parse_docx(path)
    assert result["revisi_rows"] == []
    tujuan = next(s for s in result["sections"] if s["heading"] == "1. Tujuan")
    assert "<table" in tujuan["html"]
    assert "Budi" in tujuan["html"]


def test_parse_docx_table_with_four_cols_wrong_headers_is_not_revisi(tmp_path):
    path = str(tmp_path / "wrong-headers.docx")
    document = docx.Document()
    document.add_heading("SOP Tabel Header Salah", level=0)
    document.add_heading("1. Tujuan", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Versi", "Penulis", "Isi", "Status"
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text, row[3].text = "2.0", "Ani", "Draft", "Selesai"
    document.save(path)

    result = parse_docx(path)
    assert result["revisi_rows"] == []
    tujuan = next(s for s in result["sections"] if s["heading"] == "1. Tujuan")
    assert "<table" in tujuan["html"]
    assert "Ani" in tujuan["html"]
