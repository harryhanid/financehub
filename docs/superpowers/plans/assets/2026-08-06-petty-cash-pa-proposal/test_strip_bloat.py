import os
import sys

sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
import build_deck
import strip_bloat

SRC = os.path.join(os.path.dirname(__file__), "_test_src.pptx")
DST = os.path.join(os.path.dirname(__file__), "_test_stripped.pptx")


def test_strip_shrinks_file_and_preserves_content():
    build_deck.build(SRC)
    original_size = os.path.getsize(SRC)

    strip_bloat.strip(SRC, DST)

    stripped_size = os.path.getsize(DST)
    assert stripped_size < original_size * 0.6

    prs = Presentation(DST)
    assert len(prs.slides) == 9
    assert prs.slides[0].shapes.title.text == (
        "Digitalisasi Petty Cash & PA di Bawah Rp 10 Juta via PR Portal"
    )
