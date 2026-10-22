import io
import openpyxl
from modules.sponsorship.service import seed_perusahaan, create_transaksi
from modules.sponsorship.exports import export_report_excel

SMT = 1


def _seed():
    seed_perusahaan([("SMART", "AGRI", "AGRI")])


def test_export_report_excel_has_expected_title(app):
    _seed()
    create_transaksi(SMT, {
        "cat": "Sponsorship", "description": "Sosial",
        "tanggal_form_sponsor": "2026-01-15", "cat1": "SMART",
        "lembaga_perorangan": "Yayasan Test", "amount_request": 1000000,
    })
    xls_bytes = export_report_excel(SMT, ["2026"], "all")
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    assert ws["B1"].value == "RANGKUMAN SPONSORSHIP - in Rp Full Amount"


def test_export_report_excel_contains_section_and_lembaga(app):
    _seed()
    create_transaksi(SMT, {
        "cat": "Sponsorship", "description": "Sosial",
        "tanggal_form_sponsor": "2026-01-15", "cat1": "SMART",
        "lembaga_perorangan": "Yayasan Test", "amount_request": 1000000,
    })
    xls_bytes = export_report_excel(SMT, ["2026"], "all")
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "SOSIAL" in values
    assert "Yayasan Test" in values
    assert "TOTAL SPONSORSHIP (ACTUAL)" in values


def test_export_report_excel_empty_data_still_generates_file(app):
    _seed()
    xls_bytes = export_report_excel(SMT, ["2026"], "all")
    wb = openpyxl.load_workbook(io.BytesIO(xls_bytes))
    ws = wb.active
    assert ws["B1"].value == "RANGKUMAN SPONSORSHIP - in Rp Full Amount"


def _tx(lem, tgl, cat4="Include"):
    create_transaksi(SMT, {
        "cat": "Sponsorship", "description": "Sosial", "tanggal_form_sponsor": tgl,
        "cat1": "SMART", "lembaga_perorangan": lem, "amount_request": 1000000, "cat4": cat4,
    })


def _values(xls_bytes):
    ws = openpyxl.load_workbook(io.BytesIO(xls_bytes)).active
    return [cell.value for row in ws.iter_rows() for cell in row]


def test_export_follows_the_month_filter_but_title_has_no_months(app):
    _seed()
    _tx("Jan Org", "2026-01-15"); _tx("Feb Org", "2026-02-10"); _tx("Mar Org", "2026-03-05")
    values = _values(export_report_excel(SMT, ["2026"], "all", ["01", "02"]))
    assert "RANGKUMAN SPONSORSHIP 2026 - in Rp Full Amount" in values
    assert not any(isinstance(v, str) and "(Jan" in v for v in values)
    assert "Jan Org" in values and "Feb Org" in values and "Mar Org" not in values
    assert "RANGKUMAN SPONSORSHIP - in Rp Full Amount" in values          # the fixed first title cell


def test_export_without_months_keeps_the_title_and_all_rows(app):
    _seed()
    _tx("Jan Org", "2026-01-15"); _tx("Mar Org", "2026-03-05")
    values = _values(export_report_excel(SMT, ["2026"], "all"))
    assert "RANGKUMAN SPONSORSHIP 2026 - in Rp Full Amount" in values
    assert "Jan Org" in values and "Mar Org" in values


def test_export_follows_cat4_mode_together_with_months(app):
    _seed()
    _tx("Jan Incl", "2026-01-15", "Include"); _tx("Jan Excl", "2026-01-16", "Exclude")
    include = _values(export_report_excel(SMT, ["2026"], "include", ["01"]))
    exclude = _values(export_report_excel(SMT, ["2026"], "exclude", ["01"]))
    assert "Jan Incl" in include and "Jan Excl" not in include
    assert "Jan Excl" in exclude and "Jan Incl" not in exclude


def test_export_with_a_month_that_has_no_data_still_generates_a_file(app):
    _seed()
    _tx("Jan Org", "2026-01-15")
    values = _values(export_report_excel(SMT, ["2026"], "all", ["12"]))
    assert "RANGKUMAN SPONSORSHIP 2026 - in Rp Full Amount" in values
    assert "Jan Org" not in values

def test_pct_allocation_row_is_red_italic_right_aligned_in_column_f_on_grey(app):
    _seed()
    _tx("Jan Org", "2026-01-15")
    ws = openpyxl.load_workbook(io.BytesIO(export_report_excel(SMT, ["2026"], "all"))).active
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 6).value == "% Allocation")
    assert ws.cell(row, 2).value in (None, "")
    for col in range(2, 12):
        c = ws.cell(row, col)
        assert c.fill.fgColor.rgb.endswith("D9D9D9")
        assert c.font.color.rgb.endswith("CC0066") and c.font.italic
        assert c.alignment.horizontal == "right"
