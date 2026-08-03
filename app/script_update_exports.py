import io
import openpyxl

path = r"C:\Financehub\app\modules\sponsorship\exports.py"
with open(path, "r", encoding="utf-8") as f:
    c = f.read()

new_code = """
def export_data_excel(company_id: int, filters: dict) -> bytes:
    from modules.sponsorship.service import list_transaksi
    rows = list_transaksi(company_id, filters)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Sponsorship"
    
    headers = [
        "ID", "Cat", "Description", "Tanggal Form", "Tanggal Payment",
        "Cat1", "Cat2", "Cat3", "Lembaga / Perorangan", "Perihal",
        "Amount Request", "AGRI", "APP", "ENERGY & FINANCE", "PROPERTY",
        "Reff", "Status", "Cat4", "COA", "CC Total", "CC", "Source"
    ]
    ws.append(headers)
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    
    for r in rows:
        row_data = [
            r.get("id"), r.get("cat"), r.get("description"), r.get("tanggal_form_sponsor"), r.get("tanggal_payment"),
            r.get("cat1"), r.get("cat2"), r.get("cat3"), r.get("lembaga_perorangan"), r.get("perihal"),
            r.get("amount_request"), r.get("agri"), r.get("app"), r.get("energy_finance"), r.get("property"),
            r.get("reff"), r.get("status"), r.get("cat4"), r.get("coa"), r.get("cc_total"), r.get("cc"), r.get("source")
        ]
        ws.append(row_data)
        
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_outstanding_excel(company_id: int, filters: dict) -> bytes:
    from modules.sponsorship.service import get_sponsorship_outstanding
    sections, grand_total = get_sponsorship_outstanding(company_id, filters)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Outstanding Sponsorship"
    
    headers = ["Kategori", "Tanggal", "Lembaga/Perorangan", "Amount Request", "Status"]
    ws.append(headers)
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    
    for section in sections:
        ws.append([section["cat2_label"]])
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=1).font = _SECTION_FONT
        ws.cell(row=row_idx, column=1).fill = _SECTION_FILL
        
        for r in section["rows"]:
            ws.append(["", r.get("tanggal_form_sponsor"), r.get("lembaga_perorangan"), r.get("amount_request"), r.get("status")])
            
        ws.append(["", "", "Subtotal", section["subtotal"], ""])
        row_idx = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = _SECTION_FONT
            cell.fill = _SECTION_FILL
            
    ws.append(["", "", "Grand Total", grand_total, ""])
    row_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = _GRANDTOTAL_FONT
        cell.fill = _GRANDTOTAL_FILL
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
"""

c += new_code
with open(path, "w", encoding="utf-8") as f:
    f.write(c)

print("done")
