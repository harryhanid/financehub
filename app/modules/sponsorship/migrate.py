"""One-time migration: historical Sponsorship/CSR Gedung Excel trackers -> finance_hub.db.

Usage: python migrate.py <path-to-cat1-3.xlsx> <path-to-database-awal.xlsx>

Run manually once. Not wired into app.py or any scheduler. Re-running re-imports
transaksi rows (no dedup key in the source data) — do not run twice against production.
"""
import sys
import sqlite3
import openpyxl
import config
from database import get_conn
from modules.sponsorship.service import seed_perusahaan

VALID_CAT = ("Sponsorship", "CSR Gedung")


def _cell_date(value):
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value)


def migrate_sponsorship_from_excel(perusahaan_xlsx_path: str, transaksi_xlsx_path: str, company_id: int = 1) -> dict:
    perusahaan_wb = openpyxl.load_workbook(perusahaan_xlsx_path, data_only=True)
    perusahaan_ws = perusahaan_wb.active
    perusahaan_rows = [
        (row[0], row[1], row[2])
        for row in perusahaan_ws.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]
    perusahaan_count = seed_perusahaan(perusahaan_rows)

    transaksi_wb = openpyxl.load_workbook(transaksi_xlsx_path, data_only=True)
    transaksi_ws = transaksi_wb.active
    header = [c.value for c in next(transaksi_ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}

    conn = get_conn()
    transaksi_count = 0
    skipped = []
    try:
        for row in transaksi_ws.iter_rows(min_row=2, values_only=True):
            cat = row[idx["Cat"]]
            if cat not in VALID_CAT:
                skipped.append(f"row with invalid Cat: {cat!r}")
                continue

            tanggal_form = row[idx["Tanggal Form Sponsor"]]
            if not hasattr(tanggal_form, "strftime"):
                year = row[idx["Year"]]
                if isinstance(year, int):
                    tanggal_form_str = f"{year}-01-01"
                else:
                    skipped.append(f"row with lembaga={row[idx['Lembaga/Perorangan']]!r}: Tanggal Form Sponsor is not a date ({tanggal_form!r}) and Year is also missing/invalid ({year!r})")
                    continue
            else:
                tanggal_form_str = _cell_date(tanggal_form)

            if not row[idx["Description"]] or not row[idx["Lembaga/Perorangan"]]:
                skipped.append(f"row with lembaga={row[idx['Lembaga/Perorangan']]!r}: missing required Description or Lembaga/Perorangan")
                continue

            try:
                conn.execute(
                    """
                    INSERT INTO sponsorship_transaksi
                        (company_id, cat, description, tanggal_form_sponsor, tanggal_payment,
                         cat1, cat2, cat3, lembaga_perorangan, perihal, amount_request, reff,
                         status, cat4, coa, cc_total, cc, source)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        company_id, cat, row[idx["Description"]],
                        tanggal_form_str, _cell_date(row[idx["Tanggal Payment"]]),
                        row[idx["Cat1"]], row[idx["Cat2"]], row[idx["Cat3"]],
                        row[idx["Lembaga/Perorangan"]], row[idx["Perihal"]],
                        float(row[idx["Amount Request"]] or 0),
                        None if row[idx["Reff"]] == "-" else row[idx["Reff"]],
                        row[idx["Status"]], row[idx["Cat4"]],
                        row[idx["COA"]], row[idx["CC Total"]], row[idx["CC"]], row[idx["Source"]],
                    ),
                )
                transaksi_count += 1
            except sqlite3.IntegrityError as e:
                skipped.append(f"row with lembaga={row[idx['Lembaga/Perorangan']]!r}: {e}")
        conn.commit()
    finally:
        conn.close()

    return {"perusahaan": perusahaan_count, "transaksi": transaksi_count, "skipped": skipped}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python migrate.py <path-to-cat1-3.xlsx> <path-to-database-awal.xlsx>")
        sys.exit(1)
    print("=" * 60)
    print("PERINGATAN: Migrasi ini menulis langsung ke database berikut:")
    print(f"  {config.DB_PATH}")
    print("Script ini HANYA untuk dijalankan SATU KALI terhadap data historis.")
    print("Menjalankannya lagi akan menduplikasi baris transaksi (tidak ada dedup key).")
    print("Pastikan sudah backup finance_hub.db sebelum melanjutkan.")
    print("=" * 60)
    result = migrate_sponsorship_from_excel(sys.argv[1], sys.argv[2])
    print(f"Perusahaan (lookup) diimport: {result['perusahaan']}")
    print(f"Transaksi diimport: {result['transaksi']}")
    if result["skipped"]:
        print(f"Baris di-skip ({len(result['skipped'])}):")
        for msg in result["skipped"]:
            print(f"  - {msg}")
