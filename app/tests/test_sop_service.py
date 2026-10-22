import docx
from modules.sop.service import (
    create_document, list_documents, get_document, record_upload,
    count_perlu_verifikasi_in_file,
)


def test_create_and_get_document():
    doc_id = create_document(modul="procurement", nama="SOP Master Vendor", kode="1")
    doc = get_document(doc_id)
    assert doc["modul"] == "procurement"
    assert doc["nama"] == "SOP Master Vendor"
    assert doc["status"] == "aktif"


def test_get_document_returns_none_for_missing_id():
    assert get_document(999999) is None


def test_list_documents_filters_by_modul():
    create_document(modul="procurement", nama="SOP A")
    create_document(modul="finance-payment", nama="SOP B")
    results = list_documents(modul="procurement")
    assert all(d["modul"] == "procurement" for d in results)
    assert any(d["nama"] == "SOP A" for d in results)
    assert not any(d["nama"] == "SOP B" for d in results)


def test_list_documents_filters_by_status():
    create_document(modul="hr-personnel", nama="SOP Belum", status="belum_dibuat")
    results = list_documents(status="belum_dibuat")
    assert all(d["status"] == "belum_dibuat" for d in results)


def test_list_documents_search_matches_nama_case_insensitive():
    create_document(modul="it-operations", nama="SOP Zoom Meeting")
    results = list_documents(search="zoom")
    assert any(d["nama"] == "SOP Zoom Meeting" for d in results)


def test_record_upload_updates_document_and_adds_history():
    doc_id = create_document(modul="procurement", nama="SOP Upload Test")
    record_upload(doc_id, "procurement/1-test.docx", 3, "tester", catatan="revisi awal")
    doc = get_document(doc_id)
    assert doc["file_path"] == "procurement/1-test.docx"
    assert doc["perlu_verifikasi_count"] == 3
    assert doc["uploaded_by"] == "tester"
    assert doc["uploaded_at"] is not None


def test_count_perlu_verifikasi_in_file(tmp_path):
    path = str(tmp_path / "contoh.docx")
    document = docx.Document()
    document.add_paragraph("Baris normal.")
    document.add_paragraph("Baris ⚠ [PERLU VERIFIKASI: ambang nominal].")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "PERLU VERIFIKASI di tabel"
    document.save(path)
    assert count_perlu_verifikasi_in_file(path) == 2


def test_record_upload_atomic_on_failure():
    """Task 9 (atomic-flows.md): UPDATE sop_documents + INSERT sop_upload_history dalam satu batch.
    Paksa INSERT history gagal (uploaded_by NULL -> NOT NULL): dokumen tidak boleh berubah dan
    history tidak bertambah."""
    import pytest
    from database import get_conn

    doc_id = create_document("procurement", "Dok Atomik", kode="7", status="belum_dibuat")
    before_doc = get_document(doc_id)
    conn = get_conn()
    before_hist = conn.execute("SELECT COUNT(*) FROM sop_upload_history").fetchone()[0]
    conn.close()

    with pytest.raises(Exception):
        record_upload(doc_id, "procurement/7-atomik.docx", 2, None, catatan="harus gagal")

    assert get_document(doc_id) == before_doc
    conn = get_conn()
    assert conn.execute("SELECT COUNT(*) FROM sop_upload_history").fetchone()[0] == before_hist
    conn.close()


def test_record_upload_uploaded_at_is_sqlite_style_text():
    """uploaded_at harus teks 'YYYY-MM-DD HH:MM:SS' seperti CURRENT_TIMESTAMP SQLite."""
    import re
    doc_id = create_document("procurement", "Dok Waktu", kode="8")
    record_upload(doc_id, "procurement/8.docx", 0, "tester")
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", get_document(doc_id)["uploaded_at"])


def test_list_documents_orders_kode_numerically_like_sqlite_cast():
    """CAST(kode AS INTEGER) SQLite: NULL dulu, teks bukan angka = 0, lalu urut angka ('10' setelah '9')."""
    create_document("urut", "K10", kode="10")
    create_document("urut", "K9", kode="9")
    create_document("urut", "KX", kode="abc")
    create_document("urut", "KNULL", kode=None)
    create_document("urut", "K2", kode="2")
    names = [d["nama"] for d in list_documents(modul="urut")]
    assert names == ["KNULL", "KX", "K2", "K9", "K10"]


def test_list_documents_kode_longer_than_bigint_does_not_break_list():
    """kode 25 digit: ::bigint melempar 'out of range' dan seluruh daftar SOP gagal; ::numeric tidak."""
    create_document("panjang", "K25", kode="1234567890123456789012345")
    create_document("panjang", "K3", kode="3")
    rows = [d["nama"] for d in list_documents(modul="panjang")]
    assert rows == ["K3", "K25"]
